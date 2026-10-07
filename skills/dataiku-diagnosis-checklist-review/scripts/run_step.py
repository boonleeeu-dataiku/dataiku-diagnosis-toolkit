#!/usr/bin/env python3
"""Run the reader's scripts for a checklist review and keep proof that they ran.

Usage:
  run_step.py run <orient|facts> <bundle_root> --manifest <stem>_run_manifest.json
  run_step.py verdicts <bundle_root> --manifest <stem>_run_manifest.json --checklist <checklist.xlsx>
  run_step.py verify <bundle_root> --manifest <stem>_run_manifest.json --checklist <review.xlsx> [--deck <deck.pptx>]

`run` executes the reader script, passes its output through unchanged, and records the step, exit status,
UTC time and a sha256 of the output in the manifest (hashes and statuses only, no bundle content). `facts`
output is also saved next to the manifest as `<stem>_facts.json`; its hash ignores the absolute-path fields
(`bundle_root`, `mirror`), so `verify` works on the same bundle at another path.

`verdicts` (run after `facts`) computes, with `verdicts.py`, the status of every checklist row that a rule
decides from the saved facts, prints them as JSON, saves them as `<stem>_verdicts.json` and records the step.
The status in the workbook must equal the verdict: code is final. It also lists rows whose id has a rule but
whose title differs, titles that match a rule under another id, and rows a rule leaves `undecided`; those rows get no verdict.

`verify` prints one PASS/FAIL line and exits non-zero on any problem: all three steps recorded with
exit status 0; the saved `<stem>_facts.json` and a fresh `facts.py` run both match the recorded hash (the saved
output is genuine and unedited); every ruled row's workbook status equals its verdict (each difference is named);
the checklist's first sheet is `Summary` and its Total equals the checklist's row count (the Summary was
written after the last edit); with `--deck`, the narrative and the deck file exist. Rows with no rule are not
judged. Python 3 stdlib plus openpyxl (already a plugin dependency).
"""
import argparse
import hashlib
import json
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import verdicts  # noqa: E402
import rules_platform  # noqa: E402,F401  (registers the rules with verdicts.RULES)
import rules_genai  # noqa: E402,F401
import rules_security  # noqa: E402,F401

READER = Path(__file__).resolve().parents[2] / "dataiku-diagnosis-reader"
READER_STEPS = {"orient": ["bash", str(READER / "scripts" / "orient.sh")],
                "facts": [sys.executable, str(READER / "scripts" / "facts.py")]}


def reader_version() -> str:
    m = re.search(r"^version:\s*(\S+)", (READER / "SKILL.md").read_text(encoding="utf-8"), re.M)
    return m.group(1) if m else "unknown"


FACTS_TIMEOUT = 600


def load(manifest: Path) -> dict:
    if not manifest.exists():
        return {"steps": {}}
    data = json.loads(manifest.read_text(encoding="utf-8"))
    if not isinstance(data.get("steps"), dict):
        raise ValueError("manifest has no 'steps' map")
    return data


def sibling(manifest: Path, name: str) -> Path:
    """`<stem>_<name>.json` beside the manifest, e.g. the saved facts or verdicts."""
    return manifest.with_name(re.sub(r"_run_manifest$", "", manifest.stem) + f"_{name}.json")


PATH_FIELDS = ("bundle_root", "mirror")  # facts.py prints the absolute bundle location; it is not part of the facts


def facts_digest(text: str) -> str:
    """sha256 of facts.py output without its absolute-path fields, so the same bundle at another path (a staging copy,
    another machine) hashes the same."""
    try:
        doc = json.loads(text)
        canonical = json.dumps({k: v for k, v in doc.items() if k not in PATH_FIELDS}, sort_keys=True)
    except (ValueError, AttributeError):
        canonical = text
    return hashlib.sha256(canonical.encode()).hexdigest()


def facts_match(text: str, recorded: str) -> bool:
    """Recorded hashes from before 0.32.2 were of the raw output; accept those too."""
    return recorded in (facts_digest(text), hashlib.sha256(text.encode()).hexdigest())


def record(data: dict, step: str, returncode: int, stdout: str) -> None:
    digest = facts_digest(stdout) if step == "facts" else hashlib.sha256(stdout.encode()).hexdigest()
    data["steps"][step] = {"exit_status": returncode, "at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                           "stdout_sha256": digest}


def run(step: str, bundle: str, manifest: Path) -> int:
    try:
        proc = subprocess.run([*READER_STEPS[step], bundle], capture_output=True, text=True)
        data = load(manifest)
        data["reader_version"] = reader_version()
    except (OSError, ValueError) as e:
        print(f"run_step.py: cannot run {step}: {e}", file=sys.stderr)
        return 2
    sys.stdout.write(proc.stdout)
    sys.stderr.write(proc.stderr)
    manifest.parent.mkdir(parents=True, exist_ok=True)
    record(data, step, proc.returncode, proc.stdout)
    manifest.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")
    if step == "facts" and proc.returncode == 0:
        sibling(manifest, "facts").write_text(proc.stdout, encoding="utf-8")
    return proc.returncode


def read_rows(path: Path) -> list[dict]:
    """Every checklist row (id, title, validation_status, sheet) from the section sheets, i.e. those with an `id` column."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        rows = []
        for ws in wb.worksheets:
            data = ws.iter_rows(values_only=True)
            header = list(next(data, ()) or ())
            if ws.title == "Summary" or "id" not in header:
                continue
            col = {name: header.index(name) for name in ("id", "title", "validation_status") if name in header}
            for r in data:
                if r and r[col["id"]]:
                    rows.append({"id": r[col["id"]], "sheet": ws.title,
                                 **{k: (r[i] if i < len(r) else None) for k, i in col.items() if k != "id"}})
        return rows
    finally:
        wb.close()


def run_verdicts(manifest: Path, checklist: Path) -> int:
    try:
        facts_doc = json.loads(sibling(manifest, "facts").read_text(encoding="utf-8"))
        rows = read_rows(checklist)
        data = load(manifest)
    except (OSError, ValueError) as e:
        print(f"run_step.py: cannot compute verdicts (run `facts` first, and pass a readable checklist): {e}", file=sys.stderr)
        return 2
    out = json.dumps(verdicts.compute(facts_doc, rows), indent=2) + "\n"
    sys.stdout.write(out)
    manifest.parent.mkdir(parents=True, exist_ok=True)
    sibling(manifest, "verdicts").write_text(out, encoding="utf-8")
    record(data, "verdicts", 0, out)
    manifest.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")
    return 0


def checklist_problems(path: Path) -> list[str]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception as e:  # missing, unreadable or not an xlsx
        return [f"cannot open checklist {path}: {e}"]
    try:
        if wb.sheetnames[0] != "Summary":
            return ["checklist's first sheet is not Summary (write_summary not run)"]
        total = next((row[1] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0] == "Total"), None)
        rows = sum(1 for ws in wb.worksheets[1:] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0])
    finally:
        wb.close()
    if total != rows:
        return [f"Summary Total {total} != {rows} checklist rows (Summary is stale or missing)"]
    return []


def status_problems(manifest: Path, checklist: Path, steps: dict) -> list[str]:
    """The saved facts must be the recorded run's output, and every ruled row's status must equal its verdict."""
    if "facts" not in steps:
        return []
    try:
        saved = sibling(manifest, "facts").read_text(encoding="utf-8")
    except OSError as e:
        return [f"saved facts missing: {e}"]
    if not facts_match(saved, steps["facts"]["stdout_sha256"]):
        return ["the saved facts file does not match the recorded hash (edited after the run)"]
    try:
        rows = read_rows(checklist)
    except Exception:  # an unreadable checklist is already reported by checklist_problems
        return []
    return verdicts.status_problems(verdicts.compute(json.loads(saved), rows), rows)


def verify(bundle: str, manifest: Path, checklist: Path, deck: Path | None) -> list[str]:
    problems = []
    try:
        steps = load(manifest)["steps"]
    except (OSError, ValueError) as e:
        return [f"cannot read manifest {manifest}: {e}"]
    for step in (*READER_STEPS, "verdicts"):
        if step not in steps:
            problems.append(f"{step} not recorded in the manifest")
        elif steps[step]["exit_status"] != 0:
            problems.append(f"{step} exited {steps[step]['exit_status']}")
    if "facts" in steps:
        try:
            fresh = subprocess.run([*READER_STEPS["facts"], bundle], capture_output=True, text=True, timeout=FACTS_TIMEOUT)
        except (OSError, subprocess.TimeoutExpired) as e:
            problems.append(f"fresh facts.py run failed: {e}")
        else:
            if fresh.returncode != 0:
                problems.append(f"fresh facts.py run exited {fresh.returncode}")
            elif not facts_match(fresh.stdout, steps["facts"]["stdout_sha256"]):
                problems.append("a fresh facts.py run does not match the recorded hash")
    problems += checklist_problems(checklist)
    problems += status_problems(manifest, checklist, steps)
    if deck is not None:
        narrative = checklist.with_name(checklist.stem + "_narrative.json")
        problems += [f"{p} missing" for p in (narrative, deck) if not p.exists()]
    return problems


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("step", choices=sorted(READER_STEPS))
    d = sub.add_parser("verdicts")
    d.add_argument("--checklist", required=True)
    v = sub.add_parser("verify")
    v.add_argument("--checklist", required=True)
    v.add_argument("--deck")
    for p in (r, d, v):
        p.add_argument("bundle_root")
        p.add_argument("--manifest", required=True)
    a = ap.parse_args()
    if not Path(a.bundle_root).is_dir():
        sys.exit(f"run_step.py: bundle_root not found: {a.bundle_root}")
    manifest = Path(a.manifest)
    if a.cmd == "run":
        return run(a.step, a.bundle_root, manifest)
    if a.cmd == "verdicts":
        return run_verdicts(manifest, Path(a.checklist))
    problems = verify(a.bundle_root, manifest, Path(a.checklist), Path(a.deck) if a.deck else None)
    print("RUN VERIFY: PASS" if not problems else "RUN VERIFY: FAIL - " + "; ".join(problems))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
