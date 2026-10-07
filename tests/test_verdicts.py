"""verdicts.py (deterministic statuses) and its enforcement in run_step.py verify. Rules are injected per test:
Batch 0 ships the framework with no real rules, so these pin the contract the rules will rely on."""
import json
import subprocess
import sys

import pytest

from conftest import FIXTURES, REPO_ROOT

SCRIPTS = REPO_ROOT / "skills" / "dataiku-diagnosis-checklist-review" / "scripts"
sys.path.insert(0, str(SCRIPTS))
import run_step  # noqa: E402
import verdicts  # noqa: E402

BUNDLE = FIXTURES / "bundles" / "synthetic_design_baseline"
FACTS = {"facts": {"node": {"value": {"installid": "X1"}, "source": "install.ini"}, "cgroups": {"value": "ABSENT"}}}


@pytest.fixture
def rules(monkeypatch):
    """An isolated rule registry with one rule: A-1 'First check' is Fail unless an install id exists."""
    registry = {}
    monkeypatch.setattr(verdicts, "RULES", registry)

    @verdicts.rule("A-1", "First check")
    def a1(facts):
        installid = verdicts.fact(facts, "node")["installid"]
        return verdicts.verdict("Pass" if installid != "ABSENT" else "Fail", f"installid={installid}", installid=installid)

    return registry


def rows(*pairs):
    return [{"id": i, "title": t, "validation_status": s} for i, t, s in pairs]


def test_a_ruled_row_gets_a_verdict_with_status_values_and_reason(rules):
    doc = verdicts.compute(FACTS, rows(("A-1", "First check", None), ("Z-9", "Unruled", None)))
    assert doc["verdicts"] == [{"id": "A-1", "title": "First check", "status": "Pass", "reason": "installid=X1",
                                "deciding_values": {"installid": "X1"}}]
    assert doc["rows"] == 2 and doc["rules"] == 1 and not doc["title_mismatch"] and not doc["probable_renumber"]


def test_title_match_ignores_case_punctuation_and_spacing(rules):
    assert verdicts.compute(FACTS, rows(("A-1", "  FIRST   check! ", None)))["verdicts"]


def test_same_id_with_a_different_title_gets_no_verdict_and_is_reported(rules):
    doc = verdicts.compute(FACTS, rows(("A-1", "Something else entirely", None)))
    assert doc["verdicts"] == [] and doc["title_mismatch"] == [
        {"id": "A-1", "row_title": "Something else entirely", "rule_title": "First check"}]


def test_a_renumbered_row_is_reported_not_judged(rules):
    doc = verdicts.compute(FACTS, rows(("B-7", "First check", None)))
    assert doc["verdicts"] == [] and doc["probable_renumber"] == [{"row_id": "B-7", "title": "First check", "rule_id": "A-1"}]


def test_absent_fact_is_the_string_absent_never_missing_or_none():
    assert verdicts.fact(FACTS["facts"], "cgroups") == "ABSENT"
    assert verdicts.fact(FACTS["facts"], "no_such_fact") == "ABSENT"


@pytest.mark.parametrize("bad", [
    {"status": "Passed", "reason": "x", "deciding_values": {}},          # not a fixed status
    {"status": "Pass", "reason": "two\nlines", "deciding_values": {}},    # reason must be one line
    {"status": "Pass", "reason": " ", "deciding_values": {}},
    {"status": "Pass", "reason": "ok", "deciding_values": "x"},
])
def test_a_malformed_rule_result_is_rejected(rules, bad):
    rules["A-1"] = ("First check", lambda facts: bad)
    with pytest.raises(ValueError):
        verdicts.compute(FACTS, rows(("A-1", "First check", None)))


def test_duplicate_rule_registration_is_rejected(rules):
    with pytest.raises(ValueError):
        verdicts.rule("A-1", "again")(lambda f: None)


def test_status_problems_names_each_differing_ruled_row_and_skips_unruled(rules):
    doc = verdicts.compute(FACTS, rows(("A-1", "First check", None)))
    assert verdicts.status_problems(doc, rows(("A-1", "First check", "Pass"), ("Z-9", "x", "Fail"))) == []
    out = verdicts.status_problems(doc, rows(("A-1", "First check", "Fail")))
    assert out == ["A-1: workbook Fail, verdict Pass (installid=X1)"]
    assert "blank" in verdicts.status_problems(doc, rows(("A-1", "First check", None)))[0]
    assert "not in the workbook" in verdicts.status_problems(doc, [])[0]


# --- enforcement through run_step.verify -------------------------------------------------------------

def _sh(*args):
    return subprocess.run([sys.executable, str(SCRIPTS / "run_step.py"), *map(str, args)], capture_output=True, text=True)


def _recorded(tmp_path, status):
    import openpyxl

    xlsx, manifest = tmp_path / "r.xlsx", tmp_path / "r_run_manifest.json"
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb["Summary"].append(["Total", 1])
    ws = wb.create_sheet("Section")
    ws.append(["id", "title", "validation_status"])
    ws.append(["A-1", "First check", status])
    wb.save(xlsx)
    for step in ("orient", "facts"):
        assert _sh("run", step, BUNDLE, "--manifest", manifest).returncode == 0
    assert _sh("verdicts", BUNDLE, "--manifest", manifest, "--checklist", xlsx).returncode == 0
    return manifest, xlsx


def test_verify_passes_when_workbook_statuses_equal_the_verdicts(rules, tmp_path):
    manifest, xlsx = _recorded(tmp_path, "Pass")  # the baseline bundle has an install id, so the rule says Pass
    assert run_step.verify(str(BUNDLE), manifest, xlsx, None) == []


def test_verify_fails_naming_the_item_when_the_model_deviates(rules, tmp_path):
    manifest, xlsx = _recorded(tmp_path, "Fail")
    assert run_step.verify(str(BUNDLE), manifest, xlsx, None) == ["A-1: workbook Fail, verdict Pass (installid=SYNTHETICINSTALL01)"]


def test_verify_fails_when_the_saved_facts_were_edited_after_the_run(rules, tmp_path):
    manifest, xlsx = _recorded(tmp_path, "Pass")
    facts_file = tmp_path / "r_facts.json"
    doc = json.loads(facts_file.read_text())
    doc["facts"]["node"]["value"]["installid"] = "ABSENT"
    facts_file.write_text(json.dumps(doc))
    assert any("edited after the run" in p for p in run_step.verify(str(BUNDLE), manifest, xlsx, None))


def test_verdicts_command_needs_the_facts_step_first(tmp_path):
    import openpyxl

    xlsx = tmp_path / "c.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["id", "title"])
    wb.save(xlsx)
    out = _sh("verdicts", BUNDLE, "--manifest", tmp_path / "r_run_manifest.json", "--checklist", xlsx)
    assert out.returncode == 2 and "run `facts` first" in out.stderr and "Traceback" not in out.stderr
