"""The completed-review checker (tests/lib/check_review_output.py) and the eval fixtures
it scores against. A hand-built "reference answer" workbook for each scenario must pass
cleanly; deliberately broken copies must not."""

import shutil
from pathlib import Path

import openpyxl
import pytest
import yaml

import check_review_output as cro
import read_checklist
from conftest import FIXTURES

SCENARIOS = sorted(p.stem for p in (FIXTURES / "expected").glob("*.yaml"))
EVAL_CHECKLIST = FIXTURES / "checklists" / "eval_checklist.xlsx"


def eval_items():
    wb = openpyxl.load_workbook(EVAL_CHECKLIST)
    return {ws.title: read_checklist.parse_section_sheet(ws, ws.title) for ws in wb.worksheets}


def reference_answer(rg_builder, scenario: str, path: Path, **overrides) -> Path:
    """A workbook a perfect review of `scenario` could produce: the first allowed status
    and every required mention, citing a file that exists in the bundle."""
    spec = yaml.safe_load((FIXTURES / "expected" / f"{scenario}.yaml").read_text())
    sections = {}
    for sheet, items in eval_items().items():
        sections[sheet] = []
        for it in items:
            rules = {**spec["items"][it.id], **overrides.get(it.id, {})}
            words = list(rules.get("mentions", [])) + list(rules.get("mentions_any", [])[:1])
            notes = "Currency could not be verified against Dataiku's release information." if rules.get("matches") else ""
            sections[sheet].append(rg_builder.item(
                it.id, rules["status"][0], priority=it.priority, title=it.title,
                evidence=rules.get("evidence", "data_dataiku/design/install.ini: " + "; ".join(map(str, words))),
                notes=notes,
            ))
    return rg_builder.build_checklist(path, sections)


def run(path, scenario):
    return cro.check(path, FIXTURES / "bundles" / scenario, FIXTURES / "expected" / f"{scenario}.yaml")


def test_every_scenario_has_a_bundle():
    assert SCENARIOS
    for s in SCENARIOS:
        spec = yaml.safe_load((FIXTURES / "expected" / f"{s}.yaml").read_text())
        assert (FIXTURES / spec["bundle"]).is_dir(), s
        assert (FIXTURES / "bundles" / s).is_dir()


@pytest.mark.parametrize("scenario", SCENARIOS)
def test_expected_items_match_the_eval_checklist(scenario):
    spec = yaml.safe_load((FIXTURES / "expected" / f"{scenario}.yaml").read_text())
    checklist_ids = {it.id for items in eval_items().values() for it in items}
    assert set(spec["items"]) == checklist_ids
    for item_id, rules in spec["items"].items():
        assert set(rules["status"]) <= set(cro.STATUSES), item_id


def test_eval_checklist_is_a_subset_of_the_bundled_template():
    from conftest import TEMPLATE

    template = {it.id: it for ws in openpyxl.load_workbook(TEMPLATE).worksheets
                for it in read_checklist.parse_section_sheet(ws, ws.title)}
    for items in eval_items().values():
        for it in items:
            assert template[it.id].statement == it.statement, f"{it.id} drifted from the template"


def test_fixture_bundles_are_marked_synthetic_and_not_gitignored_names():
    for bundle in (FIXTURES / "bundles").iterdir():
        assert not bundle.name.startswith("dku_diagnosis_"), "this prefix is gitignored"
        assert "SYNTHETIC TEST FIXTURE" in (bundle / "diag.txt").read_text()


@pytest.mark.parametrize("scenario", SCENARIOS)
def test_reference_answer_passes(rg_builder, tmp_path, scenario):
    report = run(reference_answer(rg_builder, scenario, tmp_path / "ref.xlsx"), scenario)
    assert report.structure == []
    assert report.citations == []
    assert [r for r in report.items.values() if not r.ok] == []
    assert report.passed


def test_wrong_status_is_scored_as_a_failure(rg_builder, tmp_path):
    path = reference_answer(rg_builder, "synthetic_design_baseline", tmp_path / "x.xlsx",
                            **{"SEC-006": {"status": ["Fail"]}})
    report = run(path, "synthetic_design_baseline")
    assert not report.passed
    assert report.items["SEC-006"].failures == ["status 'Fail' not in ['Needs Review']"]


def test_missing_mention_is_scored_as_a_failure(rg_builder, tmp_path):
    path = reference_answer(rg_builder, "synthetic_design_baseline", tmp_path / "x.xlsx",
                            **{"SEC-006": {"mentions": []}})
    report = run(path, "synthetic_design_baseline")
    assert report.items["SEC-006"].failures == ["does not mention 'proxy'"]


def test_invented_evidence_path_is_flagged(rg_builder, tmp_path):
    path = reference_answer(rg_builder, "synthetic_design_baseline", tmp_path / "x.xlsx",
                            **{"ARCH-004": {"evidence": "run/sanity-check.json and config/storage-settings.json"}})
    report = run(path, "synthetic_design_baseline")
    assert report.citations == ["ARCH-004: cites 'config/storage-settings.json', which is not in the bundle."]


def test_unknown_status_and_missing_evidence_are_structural(rg_builder, tmp_path):
    sections = {"Architecture, Compute & Infrast": [
        rg_builder.item("ARCH-001", "Passed", evidence="x"),
        rg_builder.item("ARCH-002", "Fail", evidence=""),
    ]}
    path = rg_builder.build_checklist(tmp_path / "x.xlsx", sections)
    report = cro.check(path)
    assert any("'Passed' is not one of" in p for p in report.structure)
    assert "ARCH-002: Fail with empty evidence_found." in report.structure


def test_summary_not_first_is_structural(rg_builder, tmp_path):
    path = rg_builder.build_checklist(tmp_path / "x.xlsx")
    wb = openpyxl.load_workbook(path)
    wb.move_sheet("Summary", offset=1)
    wb.save(path)
    assert any("not the first sheet" in p for p in cro.check(path).structure)


def test_data_warnings_are_structural(rg_builder, tmp_path):
    path = rg_builder.build_checklist(tmp_path / "x.xlsx", total_override=99)
    assert any(p.startswith("data_warning: ") for p in cro.check(path).structure)


def test_cited_path_extraction():
    text = ("apps/dss/design/config/general-settings.json -> cgroupSettings.enabled=false; "
            "see run/hs_err_pid*.log, <mirror>/config/x.json, https://doc.dataiku.com/a/b.html "
            "and config/project-deployer/ (2 infras); no config/api-deployer/ directory; "
            "no Spark/container configs; ERROR entries 2026/01/01-2026/01/05 at 02:00")
    assert cro.cited_paths(text) == ["apps/dss/design/config/general-settings.json"]


def test_cli_exit_codes(rg_builder, tmp_path):
    good = reference_answer(rg_builder, "synthetic_design_k8s_remote", tmp_path / "good.xlsx")
    args = ["--bundle", str(FIXTURES / "bundles" / "synthetic_design_k8s_remote"),
            "--expected", str(FIXTURES / "expected" / "synthetic_design_k8s_remote.yaml")]
    assert cro.main([str(good), *args]) == 0
    bad = tmp_path / "bad.xlsx"
    shutil.copy(good, bad)
    wb = openpyxl.load_workbook(bad)
    del wb["Summary"]
    wb.save(bad)
    assert cro.main([str(bad), *args]) == 1
