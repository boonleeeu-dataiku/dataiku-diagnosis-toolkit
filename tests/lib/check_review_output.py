#!/usr/bin/env python3
"""Check a completed checklist-review workbook (the output of the
dataiku-diagnosis-checklist-review skill) for structural soundness, and
optionally score it against an expected-answers YAML.

    python tests/lib/check_review_output.py REVIEW.xlsx [--bundle DIR] [--expected YAML] [--json]

Checks, in three groups:
  structure  Summary sheet first; every item sheet keeps the 26-column schema the deck
             generator needs; every item has a status from the fixed 5-value vocabulary,
             evidence for Pass/Fail/Partial, and a validated_at; and the review generator's
             own data-consistency checks (collect_data_warnings) find nothing.
  citations  (--bundle) every file path (with an extension) cited in evidence_found exists
             in the bundle, so invented evidence is caught.
  expected   (--expected) per-item allowed statuses and required mentions; see
             tests/fixtures/expected/*.yaml for the format.

Exit code 0 only if there are no structure or citation problems and every expected item
matches. Reuses mcp-server-review-generator's read_checklist / build_deck rather than
re-implementing the schema.
"""

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
_RG_SCRIPTS = REPO_ROOT / "mcp-server-review-generator" / "scripts"
if str(_RG_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_RG_SCRIPTS))

import openpyxl  # noqa: E402
import yaml  # noqa: E402

import build_deck  # noqa: E402
import data_checks  # noqa: E402
import read_checklist  # noqa: E402
import write_summary  # noqa: E402

STATUSES = build_deck.STATUS_ORDER
EVIDENCE_REQUIRED = {"Pass", "Fail", "Partial"}

# A path with at least one directory and a known file extension. Bare directories are
# deliberately not checked: reviews routinely cite them to say they're absent ("no
# config/api-deployer/ directory"), and "Spark/K8s" or "2026/01/05" would read as paths too.
# Placeholders/globs (<node>, *, ${...}) are skipped later.
_PATH_RE = re.compile(
    r"(?<![\w/.:-])((?:[\w.${}<>*-]+/)+[\w.${}<>*-]+\.(?:json|ini|log|txt|xml|ya?ml|conf|db|sh|py)\b)"
)


@dataclass
class ItemResult:
    id: str
    status: str
    ok: bool = True
    failures: list = field(default_factory=list)


@dataclass
class Report:
    workbook: str
    structure: list = field(default_factory=list)
    citations: list = field(default_factory=list)
    items: dict = field(default_factory=dict)
    expected_failures: int = 0

    @property
    def passed(self) -> bool:
        return not self.structure and not self.citations and self.expected_failures == 0

    def to_dict(self) -> dict:
        d = asdict(self)
        d["passed"] = self.passed
        return d


def _bundle_files(bundle_root: Path) -> list[str]:
    return [p.relative_to(bundle_root).as_posix() for p in bundle_root.rglob("*")]


def _citation_exists(cited: str, bundle_paths: list[str]) -> bool:
    """A cited path counts as real if some bundle path ends with it. Agents cite paths
    relative to the bundle root, the data-dir mirror, or as absolute DSS paths, so an
    exact match isn't required."""
    cited = cited.strip("./").rstrip("/")
    return any(p == cited or p.endswith("/" + cited) for p in bundle_paths)


def cited_paths(text: str) -> list[str]:
    out = []
    for m in _PATH_RE.finditer(text or ""):
        p = m.group(1)
        if any(tok in p for tok in ("*", "<", "${", "http")) or p.startswith(("e.g", "i.e")):
            continue
        out.append(p)
    return out


NOTES_MAX_CHARS = 350  # SKILL.md asks for ~320; a little slack so the lint isn't brittle
NOTES_MAX_BULLETS = 4  # SKILL.md asks for 3 under the headline; lint only flags clear overruns


def notes_problems(notes: str) -> list[str]:
    """Ways `notes` breaks the slide-ready format in the checklist-review SKILL.md
    ("Format of `notes`"): too long, too many bullets, or file paths (which belong in
    evidence_found)."""
    notes = (notes or "").strip()
    problems = []
    if len(notes) > NOTES_MAX_CHARS:
        problems.append(f"is {len(notes)} chars (max ~{NOTES_MAX_CHARS})")
    bullets = sum(1 for line in notes.splitlines() if line.lstrip().startswith("•"))
    if bullets > NOTES_MAX_BULLETS:
        problems.append(f"has {bullets} bullets (max {NOTES_MAX_BULLETS})")
    if cited_paths(notes):
        problems.append("cites file paths (they belong in evidence_found)")
    return problems


def check(workbook: Path, bundle: Path | None = None, expected: Path | None = None) -> Report:
    report = Report(workbook=str(workbook))
    wb = openpyxl.load_workbook(workbook, data_only=True)

    if "Summary" not in wb.sheetnames:
        report.structure.append("No 'Summary' sheet.")
    elif wb.sheetnames[0] != "Summary":
        report.structure.append(f"'Summary' is not the first sheet (order: {wb.sheetnames}).")

    items = []
    for name in wb.sheetnames:
        if name == "Summary":
            continue
        try:
            items.extend(read_checklist.parse_section_sheet(wb[name], name))
        except read_checklist.ChecklistFormatError as e:
            report.structure.append(str(e))

    validated_at = {}
    for name in wb.sheetnames:
        if name == "Summary":
            continue
        ws = wb[name]
        header = [c.value for c in ws[1]]
        if "validated_at" in header:
            col = header.index("validated_at")
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:
                    validated_at[str(row[0]).strip()] = row[col]

    seen = set()
    for it in items:
        if it.id in seen:
            report.structure.append(f"{it.id}: duplicate ID.")
        seen.add(it.id)
        if it.validation_status not in STATUSES:
            report.structure.append(
                f"{it.id}: validation_status {it.validation_status!r} is not one of {STATUSES}."
            )
        if it.validation_status in EVIDENCE_REQUIRED and not it.evidence_found:
            report.structure.append(f"{it.id}: {it.validation_status} with empty evidence_found.")
        for problem in notes_problems(it.notes):
            report.structure.append(f"{it.id}: notes {problem}.")
        if not validated_at.get(it.id):
            report.structure.append(f"{it.id}: validated_at is empty.")
        report.items[it.id] = ItemResult(id=it.id, status=it.validation_status)

    if "Summary" in wb.sheetnames:
        try:
            summary = read_checklist.parse_summary_sheet(
                wb["Summary"], build_deck.common.load_config("section_names.yaml"))
            if not any(k in summary["overall_counts"] for k in STATUSES):
                report.structure.append(
                    "Summary 'Overall Status Counts' has no numeric status counts the deck generator can read "
                    "(formulas without cached values read as empty)."
                )
            if not summary["per_section_breakdown"]:
                report.structure.append("Summary has no 'Per-Section Breakdown' block the deck generator recognises.")
            if "Report generated" not in summary["bundle_metadata"]:
                report.structure.append("Summary has no 'Report generated:' metadata row.")
            headers = [text for _, text in read_checklist.find_section_blocks(wb["Summary"])]
            if headers != list(write_summary.HEADERS.values()):
                report.structure.append(
                    f"Summary block headers read as {headers}, expected {list(write_summary.HEADERS.values())} "
                    "(a bold cell in column A outside the five headers can split a block)."
                )
            tally = {s: sum(1 for it in items if it.validation_status == s) for s in STATUSES}
            tally["Total"] = len(items)
            if summary["overall_counts"] != tally:
                report.structure.append(
                    f"Summary Overall Status Counts {summary['overall_counts']} != section-sheet tally {tally}."
                )
            for block, wanted in (("critical_findings", {"Fail"}), ("other_must_have", {"Partial", "Needs Review"})):
                expected_ids = {it.id for it in items if it.is_must_have and it.validation_status in wanted}
                got_ids = {r.get("id") for r in summary[block]}
                if got_ids != expected_ids:
                    report.structure.append(
                        f"Summary {block} lists {sorted(map(str, got_ids))}, expected {sorted(expected_ids)}."
                    )
        except read_checklist.ChecklistFormatError:
            report.structure.append("checklist format could not be parsed for the Summary comparison")
        try:
            for w in data_checks.collect_data_warnings(workbook):
                report.structure.append(f"data_warning: {w}")
        except read_checklist.ChecklistFormatError as e:
            report.structure.append(f"Summary sheet unreadable by the deck generator: {e}")

    if bundle is not None:
        paths = _bundle_files(bundle)
        for it in items:
            for cited in cited_paths(it.evidence_found):
                if not _citation_exists(cited, paths):
                    report.citations.append(f"{it.id}: cites {cited!r}, which is not in the bundle.")

    if expected is not None:
        spec = yaml.safe_load(expected.read_text(encoding="utf-8"))
        by_id = {it.id: it for it in items}
        for item_id, rules in spec["items"].items():
            it = by_id.get(item_id)
            result = report.items.setdefault(item_id, ItemResult(id=item_id, status=""))
            if it is None:
                result.failures.append("missing from workbook")
            else:
                text = f"{it.evidence_found}\n{it.notes}".lower()
                if it.validation_status not in rules.get("status", STATUSES):
                    result.failures.append(f"status {it.validation_status!r} not in {rules['status']}")
                for s in rules.get("mentions", []):
                    if str(s).lower() not in text:
                        result.failures.append(f"does not mention {s!r}")
                anys = rules.get("mentions_any", [])
                if anys and not any(str(s).lower() in text for s in anys):
                    result.failures.append(f"mentions none of {anys}")
                if rules.get("matches") and not re.search(rules["matches"], text, re.IGNORECASE):
                    result.failures.append(f"does not match /{rules['matches']}/")
            result.ok = not result.failures
        report.expected_failures = sum(1 for i in spec["items"] if not report.items[i].ok)
    return report


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--bundle", type=Path, help="bundle root, to verify cited evidence paths exist")
    ap.add_argument("--expected", type=Path, help="expected-answers YAML to score against")
    ap.add_argument("--json", action="store_true", help="print the full report as JSON")
    args = ap.parse_args(argv)

    report = check(args.workbook, args.bundle, args.expected)
    if args.json:
        print(json.dumps(report.to_dict(), indent=2))
    else:
        for group in ("structure", "citations"):
            for p in getattr(report, group):
                print(f"[{group}] {p}")
        if args.expected:
            for r in report.items.values():
                mark = "ok  " if r.ok else "FAIL"
                print(f"[expected] {mark} {r.id:<10} {r.status:<15} {'; '.join(r.failures)}")
            total = len(yaml.safe_load(args.expected.read_text())["items"])
            print(f"expected: {total - report.expected_failures}/{total} items correct")
        print("PASSED" if report.passed else "FAILED")
    return 0 if report.passed else 1


if __name__ == "__main__":
    sys.exit(main())
