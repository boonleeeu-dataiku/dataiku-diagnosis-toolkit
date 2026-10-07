#!/usr/bin/env python3
"""Regenerate the synthetic eval fixtures under tests/fixtures/.

    mcp-server-review-generator/.venv/bin/python tests/fixtures/build_fixtures.py

Writes:
  bundles/<scenario>/         hand-designed fake diagnosis bundles (see README.md)
  checklists/eval_checklist.xlsx
                              a trimmed copy of the bundled default checklist template,
                              keeping only the rows listed in EVAL_ITEM_IDS
The expected answers per scenario live in expected/<scenario>.yaml and are maintained
by hand -- they encode the checklist-review skill's calibrations, not this script.

Everything here is synthetic. Never copy real bundle data into this directory.
"""

import json
import shutil
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent.parent
TEMPLATE = REPO_ROOT / "skills" / "dataiku-diagnosis-checklist-review" / "resources" / "checklist_template.xlsx"

# Each item exercises one calibration or evidence source in the checklist-review skill.
EVAL_ITEM_IDS = [
    "ARCH-001",   # automation-node existence from a design bundle -> always Needs Review
    "ARCH-002",   # DSS version currency (web lookup, or Needs Review when unavailable)
    "ARCH-004",   # SSD storage -> sanity-check flag is Fail; else the data dir's disk ROTA from the lsblk blocks
    "ARCH-010",   # containerized exec config -> Kubernetes-conditional
    "ARCH-013",   # cluster config -> Kubernetes-conditional
    "SEC-004",    # cgroups memory limit
    "SEC-006",    # HTTPS -> reverse-proxy calibration
    "SCALE-007",  # backend.log error review -> any ERROR/WARN is Needs Review, clean is Pass
    "SCALE-008",  # backend Xmx sizing
    "SCALE-009",  # flow limits -> causal chain with OOM evidence
    "SCALE-011",  # filesystem_root connection
    "ARCH-008",   # Spark validation worded for Kubernetes -> Not Applicable without a cluster
    "ADVSEC-003", # session timeouts are real keys, 0 = unlimited
    "ADVSEC-006", # no custom post-logout redirect -> Not Applicable
    "GENAI-001",  # internal LLM Mesh code envs -> Pass internal, Needs Review non-internal
    "GENAI-009",  # Agent Hub deployer -> Needs Review when installed, Not Applicable when not
    "SEC-002",    # UIF enabled with rules -> Pass; disabled -> Fail
    "ADVSEC-008", # export restriction: any one_of key set -> Pass; none -> Fail
    "ADVSEC-009", # security headers: none -> Fail
    "SCALE-001",  # internal DB: PostgreSQL local -> Partial; remote -> Pass
    "SCALE-002",  # metastore matches estate -> Pass without a functional test
    "SCALE-003",  # graphics export on -> Pass; off -> Needs Review
    "SCALE-004",  # admin project cleanup: active schedule -> Pass; none -> Fail
    "ARCH-006",   # differently sized Spark configs -> Pass; none -> Fail
    "GENAI-007",  # Cobuild on a pre-gate DSS -> Needs Review (version gate beats feature gate)
    "SEC-001",    # instance id present in the bundle -> Pass; missing -> Fail
    "SEC-009",    # LDAP authorized groups: set -> Pass; empty -> Fail; LDAP off -> Not Applicable
    "ARCH-003",   # supported OS: web lookup -> listed Pass / not listed Fail; no lookup -> Needs Review
    "SEC-005",    # JEK-specific cgroup limits: none -> Pass; a JEK target -> Fail
    "GENAI-005",  # BYO LLM: inactive -> Not Applicable; active with project + main LLM -> Pass; either missing -> Fail
    "GENAI-006",  # BYO LLM model version: inactive -> Not Applicable; recommended -> Pass; <= 5.1 -> Fail
    "SCALE-006",  # sanity-check output with messages -> Pass; missing or empty -> Fail (never Partial)
    "SEC-010",    # SSO + LDAP enabled -> Pass; SSO disabled -> Fail; SSO enabled/LDAP off or settings missing -> Needs Review
    "GENAI-003",  # Trace Explorer default web app: set -> Pass; empty -> Fail; block missing -> Needs Review
    "GENAI-004",  # AI Services: terms accepted and enabled -> Pass; not accepted or settings missing -> Needs Review
    "SCALE-010",  # any non-blank default connection/format/engine preference -> Needs Review; all blank -> Pass
]

FIXTURE_MARKER = "SYNTHETIC TEST FIXTURE - not real diagnosis data."


def write(path: Path, content) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(content, (dict, list)):
        content = json.dumps(content, indent=2) + "\n"
    path.write_text(content, encoding="utf-8")


def root_files(root: Path, *, mem_total_mb: int, xmx: str, oom: bool, lsblk_rota: int | None = None) -> None:
    lsblk = "" if lsblk_rota is None else f"""> lsblk
NAME                  MAJ:MIN RM  SIZE RO TYPE MOUNTPOINT
sda                     8:0    0  1.5T  0 disk 
\u2514\u2500vgAPP-data_dataiku  253:3    0  1.5T  0 lvm  /data/dataiku
> lsblk -t
NAME                 ALIGNMENT MIN-IO OPT-IO PHY-SEC LOG-SEC ROTA SCHED       RQ-SIZE  RA WSAME
sda                          0  65536  65536    4096     512    {lsblk_rota} mq-deadline     256 128    0B
\u2514\u2500vgAPP-data_dataiku         0  65536  65536    4096     512    {lsblk_rota}                 128 128    0B
"""
    write(root / "diag.txt", f"""{FIXTURE_MARKER}
> uname -a
Linux synthetic-host 5.14.0-427.el9.x86_64 #1 SMP x86_64 GNU/Linux
> printenv
DKU_BACKEND_JAVA_OPTS=-Xmx{xmx} -XX:+UseG1GC
DIP_HOME=/data/dataiku/design
> free -m
              total        used        free      shared  buff/cache   available
Mem:          {mem_total_mb}       {mem_total_mb - 4000}        1200          10        2800        3500
> cat /proc/meminfo
MemTotal:       {mem_total_mb * 1024} kB
{lsblk}> cat /etc/redhat-release
Red Hat Enterprise Linux release 9.4 (Plow)
""")
    write(root / "timings.txt",
          "2026-01-01 00:00:00\t2026-01-01 00:00:01\tuname -a\n"
          "2026-01-01 00:00:01\t2026-01-01 00:00:02\tprintenv\n"
          "2026-01-01 00:00:02\t2026-01-01 00:00:03\tfree -m\n"
          "2026-01-01 00:00:03\t2026-01-01 00:00:04\tdmesg\n")
    dmesg = "[    0.000000] Linux version 5.14.0-427.el9.x86_64 (synthetic)\n"
    if oom:
        dmesg += ("[86012.123456] java invoked oom-killer: gfp_mask=0x140cca, order=0\n"
                  "[86012.123999] Out of memory: Killed process 4242 (java) total-vm:9000000kB, anon-rss:7900000kB\n"
                  "[90210.000001] Out of memory: Killed process 5151 (python3) total-vm:6000000kB, anon-rss:5800000kB\n")
    write(root / "dmesg.txt", dmesg)


def baseline_bundle(root: Path) -> None:
    """Design node, no Kubernetes, plain HTTP, populated local project deployer, cgroups
    disabled, unlimited flow concurrency with an OOM crash chain, filesystem_root present,
    rotational disk flagged by DSS's own sanity check, and an old (12.x) DSS version."""
    root_files(root, mem_total_mb=64000, xmx="2g", oom=True)
    m = root / "data_dataiku" / "design"
    write(m / "install.ini", """[general]
nodeid = synthetic-design-01
nodetype = design
installid = SYNTHETICINSTALL01

[server]
port = 11200

[javaopts]
backend.xmx = 2g
""")
    write(m / "dss-version.json", {"product_version": "12.6.0", "product_commitid": "synthetic", "conf_version": "12600"})
    write(m / "run" / "sanity-check.json", {"messages": [
        {"severity": "WARNING", "code": "WARN_MISC_DATADIR_ON_ROTATIONAL_DISK",
         "details": "The DSS data directory /data/dataiku/design is on a rotational (non-SSD) disk."},
        {"severity": "WARNING", "code": "WARN_SECURITY_NO_CGROUPS",
         "details": "cgroups resource control is not enabled."},
    ]})
    write(m / "run" / "hs_err_pid4242.log", """#
# There is insufficient memory for the Java Runtime Environment to continue.
# Native memory allocation (mmap) failed to map 1048576 bytes for committing reserved memory.
# Possible reasons:
#   The system is out of physical RAM or swap space
#
""")
    log = ["[2026/01/01-00:00:00.000] [main] [INFO] [dku.startup] - DSS backend starting"]
    for i in range(1, 6):
        log.append(f"[2026/01/0{i}-02:00:00.000] [jek-{i}] [ERROR] [dku.jobs.exec] - Job failed: "
                   "java.lang.OutOfMemoryError: Java heap space")
    log.append("[2026/01/05-03:00:00.000] [sched-1] [ERROR] [dku.scenarios] - Scenario run as deleted user 'old_admin'")
    write(m / "run" / "backend.log", "\n".join(log) + "\n")
    write(m / "config" / "projects" / "SALES_DEMO" / "params.json", {"projectKey": "SALES_DEMO", "owner": "analyst"})  # a project, but no admin one (SCALE-004 Fail)
    write(m / "config" / "general-settings.json", {
        "cgroupSettings": {"enabled": False, "cgroupsVersion": "CGROUPS_V2"},
        "useImplicitK8sCluster": False,
        "containerSettings": {"executionConfigs": []},
        "sparkSettings": {"executionConfigs": []},
        "deployerClientSettings": {"mode": "LOCAL"},
        "generativeAISettings": {"defaultRetrievableKnowledgeCodeEnv": "custom_rag_env",
                                 "defaultRetrievableKnowledgeContainerExecSelection": {"containerMode": "NONE"}},
        "impersonation": {"enabled": False, "userRules": [], "groupRules": []},
        "internalDatabase": {"connection": {"type": "PostgreSQL",
                                            "params": {"host": "127.0.0.1", "port": 5432, "db": "dss_design_db",
                                                       "password": "synthetic-not-a-secret"}}},
        "metastoreCatalogsSettings": {"synchronizeTo": {"flavor": "HIVESERVER2", "glueCredentialsMode": "DEFAULT"}},
        "hiveSettings": {"enabled": True},
        "ldapSettings": {"enabled": True, "authenticationEnabled": True, "authorizedGroups": []},
        "defaultDatasetCreationSettings": {"allowUploadsWithoutConnection": True,
                                           "preferedUploadConnection": "filesystem_managed"},
        "recipeEnginesPreferences": {"forbiddenEngines": [], "enginesPreferenceOrder": []},
        "graphicsExportsEnabled": False,
        "maxRunningActivities": 0,
        "maxRunningActivitiesPerJob": 0,
        "jekSettings": {"maxRunningJobs": 0},
        "security": {"sessionsMaxTotalTimeMinutes": 0, "sessionsMaxIdleTimeMinutes": 0,
                     "forceSingleSessionPerUser": False, "disableDataTableLinks": False},
    })
    write(m / "config" / "connections.json", {
        "filesystem_root": {"type": "Filesystem", "params": {"root": "/"}, "allowWrite": True,
                            "allowedGroups": ["administrators"]},
        "warehouse_pg": {"type": "PostgreSQL", "params": {"host": "pg.synthetic.example", "db": "dwh"},
                         "allowWrite": True},
    })
    write(m / "config" / "project-deployer" / "infras" / "prod-automation.json", {
        "id": "prod-automation", "stage": "Production",
        "automationNodes": [{"url": "https://automation.synthetic.example:11200"}],
    })
    write(m / "config" / "project-deployer" / "deployments" / "SALES_FORECAST-on-prod-automation.json", {
        "id": "SALES_FORECAST-on-prod-automation", "publishedProjectKey": "SALES_FORECAST",
        "infraId": "prod-automation", "bundleId": "v12",
    })
    write(root / "datadir_listing.txt", "\n".join([
        "  100  4 drwxr-x---   dataiku dataiku  4096 Jan  1 00:00 /data/dataiku/design/config",
        "  101  4 -rw-r-----   dataiku dataiku  1200 Jan  1 00:00 /data/dataiku/design/config/general-settings.json",
        "  102  4 -rw-r-----   dataiku dataiku   200 Jan  1 00:00 /data/dataiku/design/config/project-deployer/infras/prod-automation.json",
        "  103  4 -rw-r-----   dataiku dataiku   200 Jan  1 00:00 /data/dataiku/design/config/project-deployer/deployments/SALES_FORECAST-on-prod-automation.json",
    ]) + "\n")


def k8s_remote_bundle(root: Path) -> None:
    """Design node with a managed Kubernetes cluster and two differently sized container
    configs, DSS-terminated HTTPS, a REMOTE deployer, cgroups with a memory limit, bounded
    flow limits, a clean backend.log and no filesystem_root -- the "healthy" counterpart."""
    root_files(root, mem_total_mb=64000, xmx="8g", oom=False, lsblk_rota=0)
    m = root / "data_dataiku" / "design"
    write(m / "install.ini", """[general]
nodeid = synthetic-design-02
nodetype = design
installid = SYNTHETICINSTALL02

[server]
port = 443
ssl = true
ssl_certificate = /etc/dataiku/tls/dss.crt
ssl_certificate_key = /etc/dataiku/tls/dss.key

[javaopts]
backend.xmx = 8g
""")
    write(m / "dss-version.json", {"product_version": "14.4.3", "product_commitid": "synthetic", "conf_version": "14400"})
    write(m / "run" / "sanity-check.json", {"messages": []})
    write(m / "run" / "backend.log",
          "[2026/01/01-00:00:00.000] [main] [INFO] [dku.startup] - DSS backend starting\n"
          "[2026/01/01-00:00:05.000] [main] [INFO] [dku.startup] - DSS backend started\n")
    write(m / "config" / "general-settings.json", {
        "cgroupSettings": {
            "enabled": True, "cgroupsVersion": "CGROUPS_V2", "hierarchiesMountPoint": "/sys/fs/cgroup",
            "pythonRRecipes": {"targets": [{"cgroupPathTemplate": "DSS/${user}/pythonRRecipes",
                                            "limits": [{"key": "memory.max", "value": "40G"}]}]},
            # ~62.5 GiB host (60-120 tier, target 66% = ~41 GiB): 42G is 67% -> inside the Pass band
            "cgroups": [{"cgroupPathTemplate": "memory/DSS",
                         "limits": [{"key": "memory.limit_in_bytes", "value": "42G"}]}],
        },
        "useImplicitK8sCluster": False,
        "defaultK8sClusterId": "eks-main",
        "containerSettings": {
            "defaultExecutionConfig": "standard",
            "executionConfigs": [
                {"name": "standard", "type": "KUBERNETES", "kubernetesNamespace": "${namespace}",
                 "kubernetesResources": {"memRequestMB": 2048, "memLimitMB": 8192, "cpuRequest": 1, "cpuLimit": 2}},
                {"name": "large-memory", "type": "KUBERNETES", "kubernetesNamespace": "${namespace}",
                 "kubernetesResources": {"memRequestMB": 8192, "memLimitMB": 32768, "cpuRequest": 2, "cpuLimit": 8}},
            ],
        },
        "sparkSettings": {"executionConfigs": [
            {"name": "spark-standard", "conf": [{"key": "spark.executor.memory", "value": "4g"}]},
            {"name": "spark-large-memory", "conf": [{"key": "spark.executor.memory", "value": "12g"}]},
        ]},
        "impersonation": {"enabled": True, "userRules": [{"scope": "GLOBAL", "type": "IDENTITY"}],
                          "groupRules": []},
        "internalDatabase": {"connection": {"type": "PostgreSQL",
                                            "params": {"host": "pg-internal.synthetic.example", "port": 5432,
                                                       "db": "dss_design_db"}}},
        "metastoreCatalogsSettings": {"synchronizeTo": {"flavor": "DSS_INTERNAL"}},
        "graphicsExportsEnabled": True,
        "generativeAISettings": {
            "defaultRetrievableKnowledgeCodeEnv": "INTERNAL_retrieval_augmented_generation_v1",
            "defaultRetrievableKnowledgeContainerExecSelection": {"containerMode": "INHERIT"},
            "llmTraceSettings": {"traceExplorerDefaultWebApp": {"projectKey": "ADMINPROJECT", "webAppId": "synthWebApp"}},
        },
        "aiDrivenAnalyticsSettings": {"dataikuAIServicesTermsOfUseAccepted": True, "enabled": True},
        "deployerClientSettings": {"mode": "REMOTE", "nodeUrl": "https://deployer.synthetic.example:11200"},
        "ssoSettings": {"enabled": True, "protocol": "SAML"},
        "localAIServerSettings": {"referenceProjectKey": "LLM_REF", "mainLLMId": "openai:OpenAI:gpt-5.2"},
        "ldapSettings": {"enabled": True, "authenticationEnabled": True,
                         "authorizedGroups": ["dss-users", "dss-admins"]},
        "defaultDatasetCreationSettings": {"allowUploadsWithoutConnection": True, "virtualizable": False},
        "recipeEnginesPreferences": {"forbiddenEngines": [], "enginesPreferenceOrder": [],
                                     "forbiddenByRecipeType": {}, "preferenceByRecipeType": {}},
        "maxRunningActivities": 10,
        "maxRunningActivitiesPerJob": 4,
        "jekSettings": {"maxRunningJobs": 5},
        "security": {"sessionsMaxTotalTimeMinutes": 480, "sessionsMaxIdleTimeMinutes": 30,
                     "forceSingleSessionPerUser": True, "disableDataTableLinks": True},
    })
    write(m / "config" / "connections.json", {
        "warehouse_pg": {"type": "PostgreSQL", "params": {"host": "pg.synthetic.example", "db": "dwh"},
                         "allowWrite": True},
    })
    write(m / "config" / "clusters" / "eks-main.json", {
        "id": "eks-main", "type": "managed", "architecture": "KUBERNETES",
        "params": {"config": {"clusterId": "synthetic-eks"}},
    })
    write(m / "code-envs" / "desc" / "python" / "INTERNAL_retrieval_augmented_generation_v1" / "desc.json",
          {"envName": "INTERNAL_retrieval_augmented_generation_v1", "owner": "INTERNAL"})
    write(m / "config" / "dip.properties", "dku.exports.disableAllExports=true\n")
    write(m / "config" / "projects" / "ADMINPROJECT" / "params.json", {"projectKey": "ADMINPROJECT", "owner": "admin"})
    write(m / "config" / "projects" / "ADMINPROJECT" / "scenarios" / "CLEANUP.json", {
        "type": "step_based", "name": "CLEANUP", "active": True,
        "triggers": [{"type": "temporal", "active": True, "params": {"frequency": "Daily"}}],
        "params": {"steps": [{"type": "runnable", "name": "Clear Job logs"}]},
    })
    write(m / "config" / "plugins" / "agent-hub" / "settings.json", {"id": "agent-hub"})
    write(m / "config" / "projects" / "AGENT_HUB" / "params.json", {"projectKey": "AGENT_HUB", "owner": "svc_agents"})
    write(m / "config" / "projects" / "AGENT_HUB" / "web_apps" / "synthetic1.json", {"name": "Agent_Hub"})


def admin_python_bundle(root: Path) -> None:
    """The baseline bundle plus an `ADMINISTRATIONPROJECT` (no fixed admin project name) whose only
    scenario is an active, scheduled `custom_python` one: housekeeping can't be confirmed without
    reading its script, so SCALE-004 is Needs Review. It also has no instance id in install.ini
    (SEC-001 Fail), LDAP switched off (SEC-009 Not Applicable) SSO explicitly disabled (SEC-010 Fail), a JEK-specific cgroup target (SEC-005 Fail)
    no sanity-check output (SCALE-006 Fail; ARCH-004 can no longer be decided, so Needs Review) BYO LLM active with an old model and no reference project (GENAI-005 and GENAI-006 Fail), an empty Trace Explorer default (GENAI-003 Fail) and AI Services terms not accepted (GENAI-004 Needs Review)."""
    baseline_bundle(root)
    design = root / "data_dataiku" / "design"
    ini = design / "install.ini"
    ini.write_text("".join(ln for ln in ini.read_text().splitlines(True) if not ln.startswith("installid")))
    settings = json.loads((design / "config" / "general-settings.json").read_text())
    settings["ldapSettings"] = {"enabled": False, "authenticationEnabled": False, "authorizedGroups": []}
    (design / "run" / "sanity-check.json").unlink()
    settings["ssoSettings"] = {"enabled": False}
    settings["cgroupSettings"]["jobExecutionKernels"] = {"targets": [{"cgroupPathTemplate": "DSS/${user}/jek",
                                                                      "limits": [{"key": "memory.max", "value": "8G"}]}]}
    settings["localAIServerSettings"] = {"mainLLMId": "openai:OpenAI:gpt-5.1"}
    settings["generativeAISettings"]["llmTraceSettings"] = {"traceExplorerDefaultWebApp": {}}
    settings["aiDrivenAnalyticsSettings"] = {"dataikuAIServicesTermsOfUseAccepted": False, "enabled": False}
    write(design / "config" / "general-settings.json", settings)
    p = root / "data_dataiku" / "design" / "config" / "projects" / "ADMINISTRATIONPROJECT"
    write(p / "params.json", {"projectKey": "ADMINISTRATIONPROJECT", "owner": "admin"})
    write(p / "scenarios" / "NIGHTLY_TASKS.json", {
        "type": "custom_python", "name": "NIGHTLY_TASKS", "active": True,
        "triggers": [{"type": "temporal", "active": True, "params": {"frequency": "Daily"}}],
        "params": {"envSelection": {"envMode": "INHERIT"}},
    })
    write(p / "scenarios" / "NIGHTLY_TASKS.py", "# synthetic placeholder; never read by the review\n")


SCENARIOS = {
    "synthetic_design_baseline": baseline_bundle,
    "synthetic_design_k8s_remote": k8s_remote_bundle,
    "synthetic_design_admin_python": admin_python_bundle,
}


def build_eval_checklist(dest: Path) -> None:
    """Copy the bundled template and drop every row not in EVAL_ITEM_IDS (and any sheet
    left empty), so the eval checklist keeps the template's exact columns and wording."""
    wb = openpyxl.load_workbook(TEMPLATE)
    found = set()
    for ws in list(wb.worksheets):
        for row in range(ws.max_row, 1, -1):
            item_id = ws.cell(row=row, column=1).value
            if item_id in EVAL_ITEM_IDS:
                found.add(item_id)
            else:
                ws.delete_rows(row)
        if ws.max_row < 2:
            wb.remove(ws)
    missing = sorted(set(EVAL_ITEM_IDS) - found)
    if missing:
        raise SystemExit(f"EVAL_ITEM_IDS not found in the template: {missing}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def main() -> None:
    for name, builder in SCENARIOS.items():
        root = HERE / "bundles" / name
        shutil.rmtree(root, ignore_errors=True)
        builder(root)
        print(f"wrote {root.relative_to(REPO_ROOT)}")
    dest = HERE / "checklists" / "eval_checklist.xlsx"
    build_eval_checklist(dest)
    print(f"wrote {dest.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
