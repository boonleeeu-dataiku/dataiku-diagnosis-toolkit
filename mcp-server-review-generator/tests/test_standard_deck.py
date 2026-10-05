"""The standard (unbranded) deck: built with no branding template, only on request."""

import pytest

pytest.importorskip("pptx")
pytest.importorskip("mcp", reason="mcp SDK not installed")
from mcp.server.mcpserver.exceptions import ToolError  # noqa: E402
from pptx import Presentation  # noqa: E402

import mcp_server
import validate_deck
from conftest import item


def sections():
    return {
        "Architecture, Compute & Infrast": [
            item("ARCH-001", "Needs Review", notes="Confirm.\n• Action: ask infra team"),
            item("ARCH-002", "Fail", notes="Raise heap.\n• Action: raise backend heap"),
            item("ARCH-003", "Pass"),
        ],
        "Enterprise-grade Security, Perm": [item("SEC-001", "Partial"), item("SEC-002", "Not Applicable")],
    }


def build(checklist_factory, tmp_path, **kw):
    return mcp_server.build_platform_review_deck(
        checklist_path=str(checklist_factory(sections())), customer="Acme",
        output_path=str(tmp_path / "deck.pptx"), **kw,
    )


def test_missing_template_error_names_both_remedies(checklist_factory, tmp_path, monkeypatch):
    monkeypatch.setattr(mcp_server.common, "REPO_ROOT", tmp_path)
    with pytest.raises(ToolError, match=r"Base deck not found.*base_deck_path.*allow_standard_deck"):
        build(checklist_factory, tmp_path)
    # a wrong explicit path is never silently replaced by the fallback
    with pytest.raises(ToolError, match="Base deck not found"):
        build(checklist_factory, tmp_path, base_deck_path=str(tmp_path / "x.pptx"), allow_standard_deck=True)


def test_v1_has_no_standard_deck(checklist_factory, tmp_path, monkeypatch):
    monkeypatch.setattr(mcp_server.common, "REPO_ROOT", tmp_path)
    with pytest.raises(ToolError, match="style v1 has no unbranded fallback"):
        build(checklist_factory, tmp_path, style="v1", allow_standard_deck=True)


def test_standard_deck_builds_without_template(checklist_factory, tmp_path, monkeypatch):
    monkeypatch.setattr(mcp_server.common, "REPO_ROOT", tmp_path)  # no resources/ here: template absent
    res = build(checklist_factory, tmp_path, allow_standard_deck=True)
    assert res["base_deck_used"] == "standard" and res["branded"] is False
    assert res["structural_problems"] == []
    assert "not the Dataiku-branded deck" in res["manual_qa_checklist"]
    prs = Presentation(res["output_path"])
    assert len(prs.slides) == res["slide_count"] > 4
    cover = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert "ACME" in cover and any("not applied" in t for t in cover)
    assert not validate_deck.validate(res["output_path"])
