"""Pins the v2 deck's cover and end slides (shape names, text, relationship
types, notes) against tests/golden/v2_cover_end.json, so the cover/end code can
be moved between modules without silently changing rels, media or notes that
the visible-text golden test can't see.

Skips without the (gitignored) branding template. Refresh after an intended
change with: python3 -m pytest tests/test_v2_cover_end.py --update-golden
"""

import json

import pytest

pytest.importorskip("pptx")
from pptx import Presentation  # noqa: E402
from PIL import Image

import build_deck_v2
import common
from conftest import item

GOLDEN = common.REPO_ROOT / "tests" / "golden" / "v2_cover_end.json"

pytestmark = pytest.mark.skipif(
    not common.BRANDING_TEMPLATE.exists(),
    reason=f"branding template not present at {common.BRANDING_TEMPLATE}",
)


def sections():
    return {
        "Architecture, Compute & Infrast": [
            item("ARCH-001", "Needs Review", notes="Confirm with customer."),
            item("ARCH-002", "Fail", notes="Raise backend heap."),
            item("ARCH-003", "Pass"),
        ],
    }


def describe(slide):
    shapes = [
        {"name": sh.name, "type": str(sh.shape_type), "text": sh.text_frame.text if sh.has_text_frame else None}
        for sh in slide.shapes
    ]
    rels = sorted(
        [r.reltype.rsplit("/", 1)[-1], "external" if r.is_external else "internal"]
        for r in slide.part.rels.values()
    )
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None
    return {"shapes": shapes, "rels": rels, "notes": notes}


def summarize(path):
    prs = Presentation(str(path))
    return {"cover": describe(prs.slides[0]), "end": describe(prs.slides[len(prs.slides) - 1])}


def test_v2_cover_and_end_slides_are_pinned(checklist_factory, tmp_path, request):
    logo = tmp_path / "logo.png"
    Image.new("RGB", (40, 12), "#336699").save(logo)
    checklist = checklist_factory(sections())
    plain = tmp_path / "plain.pptx"
    with_logo = tmp_path / "logo.pptx"
    build_deck_v2.build_v2(checklist, "Acme Corp", plain, common.BRANDING_TEMPLATE)
    build_deck_v2.build_v2(checklist, "Acme Corp", with_logo, common.BRANDING_TEMPLATE, logo_path=logo)
    actual = {"no_logo": summarize(plain), "logo": summarize(with_logo)}

    if request.config.getoption("--update-golden"):
        GOLDEN.write_text(json.dumps(actual, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        pytest.skip(f"golden refreshed: {GOLDEN}")
    assert GOLDEN.exists(), "no golden yet; run with --update-golden"
    assert actual == json.loads(GOLDEN.read_text(encoding="utf-8"))
