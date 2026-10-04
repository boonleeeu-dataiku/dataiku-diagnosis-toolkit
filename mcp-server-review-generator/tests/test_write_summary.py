"""write_summary: the Summary sheet is written deterministically, and the deck
generator's reader recovers exactly what was meant."""

from datetime import date

import openpyxl
import pytest

import read_checklist
import write_summary
from conftest import default_sections, item

KEY_POINTS = {
    "ARCH-002": "Runtime DB is the embedded H2 instance.",
    "ARCH-001": "Sizing could not be confirmed from the bundle.",
    "SEC-001": "SSO is configured but MFA is not enforced.",
}
RECS = ["Fix ARCH-002 first.", "2) Review ARCH-001 with the customer.", "3. Enforce MFA (SEC-001)."]


def run(path, config, **overrides):
    kwargs = dict(
        reviewer="A. Reviewer", bundle="synthetic_bundle", node_version="Design node / 14.4.3",
        diagnosis_generated="2026-07-20", key_points=KEY_POINTS, recommendations=RECS,
        config=config, today=date(2026, 8, 1),
    )
    kwargs.update(overrides)
    return write_summary.write_summary(path, **kwargs)


def cells(path):
    ws = openpyxl.load_workbook(path)["Summary"]
    return [[c.value for c in row] for row in ws.iter_rows()]


@pytest.fixture
def checklist(checklist_factory):
    return checklist_factory()


def test_round_trips_through_the_deck_reader(checklist, section_config):
    result = run(checklist, section_config)
    data = read_checklist.read_checklist(checklist, section_config)

    assert data.overall_counts == {"Pass": 2, "Fail": 1, "Partial": 1, "Needs Review": 1, "Not Applicable": 1, "Total": 6}
    assert result["counts"] == data.overall_counts
    assert data.bundle_metadata["Bundle"] == "synthetic_bundle"
    assert data.bundle_metadata["Report generated"] == "2026-08-01"
    assert data.bundle_metadata["Reviewer"] == "A. Reviewer"
    assert [b["section"] for b in data.per_section_breakdown] == data.sheet_tab_names
    assert [r["id"] for r in data.critical_findings] == ["ARCH-002"]
    assert [r["id"] for r in data.other_must_have] == ["ARCH-001", "SEC-001"]  # key_points order
    assert data.critical_findings[0]["status"] == "Fail"
    assert data.critical_findings[0]["key_point"] == KEY_POINTS["ARCH-002"]
    assert data.recommendations == ["1. Fix ARCH-002 first.", "2. Review ARCH-001 with the customer.",
                                    "3. Enforce MFA (SEC-001)."]


def test_summary_is_first_sheet_and_exactly_five_block_headers_are_detected(checklist, section_config):
    run(checklist, section_config)
    wb = openpyxl.load_workbook(checklist)
    assert wb.sheetnames[0] == "Summary"
    blocks = read_checklist.find_section_blocks(wb["Summary"])
    assert [text for _, text in blocks] == list(write_summary.HEADERS.values())


def test_counts_are_literal_integers(checklist, section_config):
    run(checklist, section_config)
    ws = openpyxl.load_workbook(checklist)["Summary"]
    numeric = [c.value for row in ws.iter_rows(min_col=2, max_col=6) for c in row if c.value is not None
               and not isinstance(c.value, str)]
    assert numeric and all(type(v) is int for v in numeric)
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row)


def test_is_idempotent_and_leaves_section_sheets_alone(checklist, section_config):
    before = {n: [[c.value for c in r] for r in ws.iter_rows()]
              for n, ws in ((n, openpyxl.load_workbook(checklist)[n]) for n in openpyxl.load_workbook(checklist).sheetnames) if n != "Summary"}
    run(checklist, section_config)
    first = cells(checklist)
    run(checklist, section_config)
    assert cells(checklist) == first
    wb = openpyxl.load_workbook(checklist)
    assert wb.sheetnames.count("Summary") == 1
    after = {n: [[c.value for c in r] for r in wb[n].iter_rows()] for n in wb.sheetnames if n != "Summary"}
    assert after == before


def test_replaces_a_sloppy_hand_written_summary(checklist_factory, section_config):
    path = checklist_factory(overall={"Pass": 1}, breakdown=None, critical=None, other=None, recommendations=None)
    run(path, section_config)
    data = read_checklist.read_checklist(path, section_config)
    assert data.overall_counts["Total"] == 6 and data.per_section_breakdown and data.recommendations


def test_status_spelling_is_normalised(checklist_factory, section_config):
    sections = default_sections()
    sections["Architecture, Compute & Infrast"][0] = item("ARCH-001", "needs review")
    path = checklist_factory(sections=sections)
    run(path, section_config)
    data = read_checklist.read_checklist(path, section_config)
    assert data.overall_counts["Needs Review"] == 1
    assert data.other_must_have[0]["status"] == "Needs Review"


def test_no_findings_keeps_headers_and_writes_no_placeholder_rows(checklist_factory, section_config):
    sections = {"Architecture, Compute & Infrast": [item("ARCH-001", "Pass"), item("ARCH-002", "Pass")]}
    path = checklist_factory(sections=sections)
    run(path, section_config, key_points={})
    data = read_checklist.read_checklist(path, section_config)
    assert data.critical_findings == [] and data.other_must_have == []
    flat = [v for row in cells(path) for v in row]
    assert "None" not in flat


@pytest.mark.parametrize("override, match", [
    (dict(key_points={k: v for k, v in KEY_POINTS.items() if k != "SEC-001"}), "missing a line.*SEC-001"),
    (dict(key_points={**KEY_POINTS, "NOPE-1": "x"}), "unknown item id.*NOPE-1"),
    (dict(key_points={**KEY_POINTS, "ARCH-003": "a pass item"}), "not must-have Fail/Partial/Needs Review.*ARCH-003"),
    (dict(key_points={**KEY_POINTS, "ARCH-002": "x" * 91}), "exceeds 90"),
    (dict(key_points={**KEY_POINTS, "ARCH-002": "two\nlines"}), "single line"),
    (dict(recommendations=[]), "recommendations"),
    (dict(reviewer=" "), "reviewer"),
])
def test_bad_input_fails_loudly_and_leaves_file_untouched(checklist, section_config, override, match):
    original = checklist.read_bytes()
    with pytest.raises(write_summary.SummaryInputError, match=match):
        run(checklist, section_config, **override)
    assert checklist.read_bytes() == original
    assert not checklist.with_name(checklist.stem + ".summary-tmp" + checklist.suffix).exists()


def test_blank_validation_status_is_rejected(checklist_factory, section_config):
    sections = default_sections()
    sections["Architecture, Compute & Infrast"][2] = item("ARCH-003", "", priority="nice_to_have")
    path = checklist_factory(sections=sections)
    with pytest.raises(write_summary.SummaryInputError, match="ARCH-003"):
        run(path, section_config)


def test_tab_name_with_trailing_whitespace_round_trips(checklist_factory, section_config):
    """The bundled template's last tab is "Advanced Security Options (DSS " (trailing space)."""
    sections = {"Advanced Security Options (DSS ": [item("ADVSEC-001", "Fail"), item("ADVSEC-002", "Pass")]}
    path = checklist_factory(sections=sections)
    run(path, section_config, key_points={"ADVSEC-001": "Fails."})
    data = read_checklist.read_checklist(path, section_config)
    assert data.per_section_breakdown[0]["counts"]["Fail"] == 1
