import openpyxl
import pytest

import read_checklist
from conftest import item


def test_reads_default_checklist(checklist_factory, section_config):
    data = read_checklist.read_checklist(checklist_factory(), section_config)

    assert data.sheet_tab_names == ["Architecture, Compute & Infrast", "Enterprise-grade Security, Perm"]
    assert data.bundle_metadata["Report generated"] == "2026-07-22 10:00"
    assert data.overall_counts == {
        "Pass": 2, "Fail": 1, "Partial": 1, "Needs Review": 1, "Not Applicable": 1, "Total": 6,
    }
    assert [b["section"] for b in data.per_section_breakdown] == data.sheet_tab_names
    assert [c["id"] for c in data.critical_findings] == ["ARCH-002"]
    assert [o["id"] for o in data.other_must_have] == ["ARCH-001", "SEC-001"]
    assert data.recommendations == ["1. Fix ARCH-002 first.", "2. Review ARCH-001 with the customer."]

    arch = data.items_by_sheet["Architecture, Compute & Infrast"]
    assert [i.id for i in arch] == ["ARCH-001", "ARCH-002", "ARCH-003"]
    assert arch[1].validation_status == "Fail"
    assert arch[1].is_must_have and arch[1].is_flagged
    assert not arch[2].is_flagged  # Pass


def test_missing_summary_sheet_fails_fast(tmp_path, section_config):
    wb = openpyxl.Workbook()
    wb.active.title = "Architecture"
    path = tmp_path / "no_summary.xlsx"
    wb.save(path)
    with pytest.raises(read_checklist.ChecklistFormatError, match="no 'Summary' sheet"):
        read_checklist.read_checklist(path, section_config)


def test_summary_only_fails_fast(checklist_factory, section_config):
    path = checklist_factory(sections={})
    with pytest.raises(read_checklist.ChecklistFormatError, match="no checklist-item sheets"):
        read_checklist.read_checklist(path, section_config)


def test_missing_anchor_block_fails_fast(tmp_path, section_config):
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb.active["A1"] = "Something else entirely"
    wb.create_sheet("Architecture").append(read_checklist.REQUIRED_ITEM_COLUMNS)
    path = tmp_path / "no_anchor.xlsx"
    wb.save(path)
    with pytest.raises(read_checklist.ChecklistFormatError, match="Overall Status Counts"):
        read_checklist.read_checklist(path, section_config)


def test_misnamed_item_column_is_named_in_error(checklist_factory, section_config):
    cols = list(read_checklist.REQUIRED_ITEM_COLUMNS)
    cols[cols.index("validation_status")] = "status"
    path = checklist_factory(item_columns=cols)
    with pytest.raises(read_checklist.ChecklistFormatError, match="validation_status"):
        read_checklist.read_checklist(path, section_config)


def test_reordered_item_columns_fail_fast(checklist_factory, section_config):
    cols = list(read_checklist.REQUIRED_ITEM_COLUMNS)
    cols[0], cols[1] = cols[1], cols[0]
    path = checklist_factory(item_columns=cols)
    with pytest.raises(read_checklist.ChecklistFormatError, match="header order differs"):
        read_checklist.read_checklist(path, section_config)


def test_rows_without_id_are_skipped(checklist_factory, section_config):
    sections = {"Architecture, Compute & Infrast": [item("ARCH-001", "Pass"), item("", "Fail")]}
    data = read_checklist.read_checklist(checklist_factory(sections=sections), section_config)
    assert [i.id for i in data.items_by_sheet["Architecture, Compute & Infrast"]] == ["ARCH-001"]


def test_block_headers_found_by_style_not_position(checklist_factory):
    ws = openpyxl.load_workbook(checklist_factory())["Summary"]
    headers = [text for _, text in read_checklist.find_section_blocks(ws)]
    assert headers == [
        "Overall Status Counts",
        "Per-Section Breakdown",
        "Critical Findings - Must-Have Items Failing",
        "Other Must-Have Items: Partial / Needs Review",
        "Priority-Ordered Recommendations",
    ]


def test_omitted_blocks_parse_as_empty(checklist_factory, section_config):
    path = checklist_factory(breakdown=None, critical=None, other=None, recommendations=None)
    data = read_checklist.read_checklist(path, section_config)
    assert data.per_section_breakdown == []
    assert data.critical_findings == []
    assert data.other_must_have == []
    assert data.recommendations == []


def test_recommendations_table_layout_folds_extra_columns(checklist_factory, section_config):
    path = checklist_factory(recommendations_table=[
        ("#", "Action", "Related IDs"),
        (1, "Stop memory exhaustion", "SCALE-009, SEC-004"),
        (2, "Plan the upgrade", None),
    ])
    data = read_checklist.read_checklist(path, section_config)
    assert data.recommendations == [
        "Stop memory exhaustion (Related IDs: SCALE-009, SEC-004)",
        "Plan the upgrade",
    ]


def test_id_led_table_keys_follow_header_names(checklist_factory, section_config):
    path = checklist_factory(critical_cols=("ID", "Check Title"), critical=[{"id": "ARCH-002", "check title": "X"}])
    data = read_checklist.read_checklist(path, section_config)
    assert data.critical_findings == [{"id": "ARCH-002", "check_title": "X"}]


def test_summary_sheet_is_streamed_once_not_once_per_cell(checklist_factory, section_config):
    """A read-only worksheet re-streams from the top on every ws.cell() call, which made the Summary parse O(rows^2)."""
    wb = openpyxl.load_workbook(checklist_factory(), read_only=True)
    ws = wb["Summary"]
    streams = {"n": 0}
    real = ws.iter_rows

    def counting(*args, **kwargs):
        streams["n"] += 1
        return real(*args, **kwargs)

    ws.iter_rows = counting
    read_checklist.parse_summary_sheet(ws, section_config)
    wb.close()
    assert streams["n"] == 1
