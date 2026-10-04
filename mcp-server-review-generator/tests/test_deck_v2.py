"""v2 deck: analysis rules (pure), narrative validation, and a build smoke test
(skipped without the branding template)."""

import json

import pytest

import common
import deck_analysis as da
import narrative
import read_checklist
import section_names
from conftest import item

OWNERS = {"AI platform owner": ["ai platform"], "Infrastructure": ["infra team"], "Platform team": ["platform"]}


def notes(head, ev="", action=""):
    return "\n".join([head] + ([f"• {ev}"] if ev else []) + ([f"• Action: {action}"] if action else []))


def sections():
    return {
        "Architecture, Compute & Infrast": [
            item("ARCH-001", "Needs Review", notes=notes("Automation node unconfirmed", "", "confirm with platform team")),
            item("ARCH-002", "Needs Review", notes=notes("Backups unknown", "", "ask infra team for the schedule")),
            item("ARCH-003", "Pass"),
            item("ARCH-004", "Not Applicable"),
        ],
        "Enterprise-grade Security, Perm": [
            item("SEC-001", "Partial", notes=notes("Shared ceiling", "98 memory kills", "size limits; see SCALE-009")),
            item("SEC-009", "Pass", notes="Groups populated"),
        ],
        "Scalable Performance, Stability": [
            item("SCALE-009", "Fail", notes=notes("Uncapped jobs", "maxRunningJobs=0", "set finite cap; see SEC-001")),
            item("SCALE-010", "Fail", notes=notes("Stacks visible", "hideErrorStacks=false", "set true after confirming support flow")),
        ],
    }


@pytest.fixture
def analysis(checklist_factory, section_config):
    path = checklist_factory(sections(), recommendations=("1. Fix SEC-001 first.", "2. Tidy up SEC-009."))
    data = read_checklist.read_checklist(path, section_config)
    ordered = section_names.order_sections(data.sheet_tab_names, section_config)
    return data, da.analyze(data, ordered, {"owner_keywords": OWNERS})


def test_applicable_denominator_excludes_na(analysis):
    _, a = analysis
    assert (a.total, a.na, a.applicable, a.counts["Pass"]) == (8, 1, 7, 2)


def test_parse_notes_three_lines():
    assert da.parse_notes("Head\n• ev one\n• Action: do it") == ("Head", ["ev one"], "do it")


def test_root_cause_groups_follow_cross_references(analysis):
    _, a = analysis
    assert sorted(a.root_causes[0]) == ["SCALE-009", "SEC-001"]


def test_quick_wins_are_single_setting_fails_outside_root_cause(analysis):
    _, a = analysis
    assert [(q.id, q.setting, q.to) for q in a.quick_wins] == [("SCALE-010", "hideErrorStacks", "true")]
    assert a.remaining_fails == ["SCALE-009"]


def test_owner_split_covers_every_needs_review_item(analysis):
    _, a = analysis
    assert dict(a.owners) == {"Infrastructure": ["ARCH-002"], "Platform team": ["ARCH-001"]}


def test_recommendation_citing_a_pass_item_is_flagged(analysis):
    _, a = analysis
    assert any("SEC-009" in w and "Pass" in w for w in a.warnings)
    assert not any("SEC-001" in w for w in a.warnings)


def test_expand_ids_handles_slash_shorthand():
    assert da.expand_ids("Fix (SCALE-007/008) and SEC-004") == ["SCALE-007", "SCALE-008", "SEC-004"]


def test_narrative_rejects_pass_item_cited_as_problem(analysis):
    _, a = analysis
    with pytest.raises(narrative.NarrativeError, match="SEC-009.*Pass"):
        narrative.validate({"roadmap": {"now": [{"effort": "S", "action": "x", "ids": ["SEC-009"]}]}}, a)


def test_narrative_allows_pass_with_declared_caveat(analysis):
    _, a = analysis
    narrative.validate({"caveats": ["SEC-009"], "roadmap": {"now": [{"effort": "S", "action": "x", "ids": ["SEC-009"]}]}}, a)


def test_narrative_owner_groups_must_equal_needs_review_total(analysis):
    _, a = analysis
    bad = {"owners": {"groups": [{"owner": "X", "asks": [{"id": "ARCH-001", "ask": "a"}]}]}}
    with pytest.raises(narrative.NarrativeError, match="Needs Review total"):
        narrative.validate(bad, a)


def test_narrative_tile_figure_must_trace_to_cited_row(analysis):
    _, a = analysis
    with pytest.raises(narrative.NarrativeError, match="does not appear"):
        narrative.validate({"risk1": {"tiles": [{"number": "777", "label": "x", "id": "SEC-001"}]}}, a)
    narrative.validate({"risk1": {"tiles": [{"number": "98", "label": "x", "id": "SEC-001"}]}}, a)


def test_narrative_unknown_id_and_bad_effort(analysis):
    _, a = analysis
    with pytest.raises(narrative.NarrativeError, match="not in any checklist"):
        narrative.validate({"quick_wins": {"rows": [{"id": "NOPE-001"}]}}, a)
    with pytest.raises(narrative.NarrativeError, match="effort"):
        narrative.validate({"roadmap": {"now": [{"effort": "XL", "action": "x", "ids": []}]}}, a)


def test_narrative_wrong_types_are_reported_by_path(analysis):
    _, a = analysis
    bad = {"risk2": {"positives": ["HTTPS on", "exports off"], "table": []},
           "roadmap": {"now": [{"effort": "S", "action": ["Set Max Jobs"], "ids": ["SEC-001"]}], "next": "later"},
           "takeaways": [{"heading": "x", "body": ["y"], "tone": "good"}]}
    with pytest.raises(narrative.NarrativeError) as exc:
        narrative.validate(bad, a)
    msg = str(exc.value)
    for path in ("risk2.positives: expected a string, got a list", "roadmap.now[1].action", "roadmap.next: expected a list",
                 "takeaways[1].body"):
        assert path in msg


def test_narrative_correct_shapes_pass(analysis):
    _, a = analysis
    narrative.validate({"risk2": {"positives": "HTTPS on", "table": []},
                        "risk1": {"tiles": [{"number": 98, "label": "x", "id": "SEC-001"}], "fix": [{"text": "t", "id": ["SEC-001"]}]},
                        "roadmap": {"now": [{"effort": "S", "action": "Set Max Jobs", "ids": "SEC-001"}]}}, a)


def test_scaffold_exposes_shape(analysis, tmp_path):
    _, a = analysis
    f = tmp_path / "x_review.xlsx"; f.write_bytes(b"x")
    assert narrative.scaffold(a, f)["shape"]["risk2"]["positives"] == "str"


def test_scaffold_shape_lists_enum_values(analysis, tmp_path):
    _, a = analysis
    f = tmp_path / "x_review.xlsx"; f.write_bytes(b"x")
    shape = narrative.scaffold(a, f)["shape"]
    for field, values in ((shape["snapshot"][0]["state"], narrative.STATES),
                          (shape["takeaways"][0]["tone"], narrative.TONES),
                          (shape["roadmap"]["now"][0]["effort"], narrative.EFFORTS)):
        assert all(v in field for v in values)
    assert narrative.SHAPE["snapshot"][0]["state"] == "str"  # validation shape is unchanged


def test_narrative_load_errors(tmp_path):
    with pytest.raises(narrative.NarrativeError, match="not found"):
        narrative.load(tmp_path / "missing.json")
    bad = tmp_path / "bad.json"; bad.write_text("{nope")
    with pytest.raises(narrative.NarrativeError, match="not valid JSON"):
        narrative.load(bad)


@pytest.mark.skipif(not common.BRANDING_TEMPLATE.exists(), reason="branding template not present")
def test_v2_build_smoke(checklist_factory, tmp_path):
    from pptx import Presentation
    import build_deck_v2
    import validate_deck

    path = checklist_factory(sections())
    out = tmp_path / "v2.pptx"
    res = build_deck_v2.build_v2(path, "Acme", out, common.BRANDING_TEMPLATE)
    assert res["applicable_count"] == 7 and res["pass_count"] == 2
    assert validate_deck.validate(out) == []
    prs = Presentation(str(out))
    titles = [s.shapes.title.text_frame.text if s.shapes.title is not None else "" for s in prs.slides]
    assert titles[1].startswith("2 of 7 applicable")      # verdict, N/A excluded
    assert any(t.startswith("Appendix") for t in titles)
    assert not any("Thank You" in t or "Table of Contents" in t for t in titles)
    names = {sh.name for s in prs.slides for sh in s.shapes}
    assert {"Headline stat", "Takeaway card 1", "Detail table"} <= names
    cover_text = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "Generated by" not in cover_text


def test_narrative_default_path_and_staleness(tmp_path):
    assert narrative.default_path(tmp_path / "x_review.xlsx").name == "x_review_narrative.json"
    assert narrative.staleness_warnings({}, "abc") == []
    assert narrative.staleness_warnings({"checklist_sha256": "ABC"}, "abc") == []
    assert "different version" in narrative.staleness_warnings({"checklist_sha256": "zzz"}, "abc")[0]


def test_example_narrative_validates_against_synthetic_checklist(analysis):
    _, a = analysis
    ex = narrative.load(common.REPO_ROOT / "tests" / "fixtures" / "narrative_example.json")
    narrative.validate(ex, a)


@pytest.mark.skipif(not common.BRANDING_TEMPLATE.exists(), reason="branding template not present")
def test_v2_build_autodiscovers_narrative_beside_checklist(checklist_factory, tmp_path):
    import build_deck_v2
    path = checklist_factory(sections())
    (path.with_name(path.stem + "_narrative.json")).write_text(
        json.dumps({"verdict_title": "Auto-found verdict", "checklist_sha256": "stale"}))
    res = build_deck_v2.build_v2(path, "Acme", tmp_path / "o.pptx", common.BRANDING_TEMPLATE)
    assert res["narrative_used"].endswith("_narrative.json")
    assert any("different version" in w for w in res["warnings"])


@pytest.mark.skipif(not common.BRANDING_TEMPLATE.exists(), reason="branding template not present")
def test_build_tool_reports_files_it_read(checklist_factory, tmp_path):
    mcp_server = pytest.importorskip("mcp_server")
    path = checklist_factory(sections())
    (path.with_name(path.stem + "_narrative.json")).write_text(json.dumps({"verdict_title": "V"}))
    res = mcp_server.build_platform_review_deck(checklist_path=str(path), customer="Acme",
                                                output_path=str(tmp_path / "o.pptx"))
    assert res["checklist_path"] == str(path) and res["checklist_modified"]
    assert res["narrative_used"].endswith("_narrative.json") and res["narrative_modified"]


def test_scaffold_gives_every_fact_a_narrative_must_cover(analysis, tmp_path):
    data, a = analysis
    f = tmp_path / "x_review.xlsx"; f.write_bytes(b"checklist bytes")
    sc = narrative.scaffold(a, f)
    assert sc["checklist_sha256"] == __import__("hashlib").sha256(b"checklist bytes").hexdigest()
    assert sc["narrative_path"].endswith("x_review_narrative.json")
    assert sc["applicable"] == 7 and sc["counts"]["Fail"] == 2
    assert {r["id"] for r in sc["needs_review"]} == {"ARCH-001", "ARCH-002"}
    assert all(r["suggested_owner"] for r in sc["needs_review"])
    assert [r["id"] for r in sc["not_applicable"]] == ["ARCH-004"]
    assert {r["id"] for r in sc["open_items"]} == {"SEC-001", "SCALE-009", "SCALE-010"}
    assert sc["rules"] and "root_cause_groups" in sc and "quick_win_candidates" in sc


def test_scaffold_suggestions_pass_narrative_validation(analysis, tmp_path):
    data, a = analysis
    f = tmp_path / "x.xlsx"; f.write_bytes(b"x")
    sc = narrative.scaffold(a, f)
    owners = {}
    for r in sc["needs_review"]:
        owners.setdefault(r["suggested_owner"], []).append({"id": r["id"], "ask": "confirm"})
    narr = {"owners": {"groups": [{"owner": o, "asks": asks} for o, asks in owners.items()]},
            "na_groups": [{"heading": "N/A", "note": "n", "ids": [r["id"] for r in sc["not_applicable"]]}]}
    narrative.validate(narr, a)


@pytest.mark.skipif(not common.BRANDING_TEMPLATE.exists(), reason="branding template not present")
def test_v2_build_flags_a_missing_narrative(checklist_factory, tmp_path):
    import build_deck_v2
    path = checklist_factory(sections())
    res = build_deck_v2.build_v2(path, "Acme", tmp_path / "o.pptx", common.BRANDING_TEMPLATE)
    assert res["narrative_missing"] is True and res["narrative_used"] is None
    (path.with_name(path.stem + "_narrative.json")).write_text(json.dumps({}))
    res = build_deck_v2.build_v2(path, "Acme", tmp_path / "o2.pptx", common.BRANDING_TEMPLATE)
    assert res["narrative_missing"] is False


def test_id_list_accepts_string_or_list():
    import build_deck_v2
    assert build_deck_v2.id_list("SCALE-004") == ["SCALE-004"]
    assert build_deck_v2.id_list("SCALE-006 / SCALE-012") == ["SCALE-006", "SCALE-012"]
    assert build_deck_v2.id_list(["SCALE-004"]) == ["SCALE-004"]


@pytest.mark.skipif(not common.BRANDING_TEMPLATE.exists(), reason="branding template not present")
def test_v2_build_risk3_cards_accept_string_ids(checklist_factory, tmp_path):
    """A string `ids` used to index to its first character and fail with KeyError: 'S'."""
    import build_deck_v2
    path = checklist_factory(sections())
    (path.with_name(path.stem + "_narrative.json")).write_text(json.dumps({
        "risk3": {"title": "Open items", "cards": [
            {"heading": "Backups", "ids": "ARCH-002", "found": "Unknown", "todo": "Ask infra"}]}}))
    res = build_deck_v2.build_v2(path, "Acme", tmp_path / "o.pptx", common.BRANDING_TEMPLATE)
    assert res["narrative_used"].endswith("_narrative.json")


def test_risk_title_adds_prefix_only_when_missing():
    import build_deck_v2
    assert build_deck_v2.risk_title(1, "Memory is unsized") == "Risk 1: Memory is unsized"
    assert build_deck_v2.risk_title(2, "Risk 2: Hardening") == "Risk 2: Hardening"
    assert build_deck_v2.risk_title(3, "risk 3:  Ops") == "risk 3:  Ops"
    assert build_deck_v2.risk_title(1, "Risk 2: wrong slot") == "Risk 1: Risk 2: wrong slot"


def _analyse(checklist_factory, section_config, secs):
    data = read_checklist.read_checklist(checklist_factory(sections=secs), section_config)
    return da.analyze(data, section_names.order_sections(data.sheet_tab_names, section_config), {})


def test_unknown_or_blank_status_is_counted_as_needs_review_not_dropped(checklist_factory, section_config):
    secs = sections()
    secs["Architecture, Compute & Infrast"] += [item("ARCH-005", "Passed"), item("ARCH-006", "")]
    a = _analyse(checklist_factory, section_config, secs)
    assert a.rows["ARCH-005"].status == a.rows["ARCH-006"].status == "Needs Review"
    assert a.total == 10 and a.counts["Needs Review"] == 4


def test_duplicate_item_id_keeps_the_first_occurrence(checklist_factory, section_config):
    secs = sections()
    secs["Enterprise-grade Security, Perm"].append(item("ARCH-003", "Fail"))
    a = _analyse(checklist_factory, section_config, secs)
    assert a.rows["ARCH-003"].status == "Pass"
    assert a.total == 8


def test_build_result_reports_what_the_deck_rendered(analysis):
    import build_deck_v2
    _, a = analysis
    assert build_deck_v2.rendered_quick_wins(a, {}) == [q.id for q in a.quick_wins]
    narr = {"quick_wins": {"rows": [{"id": "SCALE-009"}]},
            "owners": {"groups": [{"owner": "Ops", "asks": [{"id": "X-1"}]}, {"owner": "Empty", "asks": []}]}}
    assert build_deck_v2.rendered_quick_wins(a, narr) == ["SCALE-009"]
    assert build_deck_v2.rendered_owner_groups(a, narr) == {"Ops": ["X-1"]}
    assert build_deck_v2.rendered_owner_groups(a, {}) == {o: ids for o, ids in a.owners}
