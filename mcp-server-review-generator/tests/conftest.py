"""Shared pytest fixtures: puts scripts/ on sys.path (it has no __init__.py and
uses bare sibling imports -- see CLAUDE.md's "Import resolution" note) and
builds synthetic checklist workbooks on the fly with openpyxl, so tests never
depend on a real (customer-derived) checklist.

The builder mirrors the layout of a checklist produced by the
dataiku-diagnosis-checklist-review skill: a Summary sheet (metadata rows, then
narrative blocks whose header cells share one style signature) followed by one
26-column sheet per review section. Status counts are written as literal
numbers, not COUNTIF formulas -- openpyxl writes formulas without cached
values, which read_checklist (data_only=True) would see as empty.
"""

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import read_checklist  # noqa: E402

STATUSES = ["Pass", "Fail", "Partial", "Needs Review", "Not Applicable"]
BLOCK_FONT = Font(bold=True, size=12)
TABLE_HEADER_FONT = Font(bold=True, size=11)
TABLE_HEADER_FILL = PatternFill("solid", fgColor="FFE8EEF7")


def item(id, status, priority="must_have", title=None, statement=None, evidence="", notes=""):
    return {
        "id": id,
        "priority": priority,
        "title": title or f"Title of {id}",
        "statement": statement or f"Statement of {id}",
        "statement_short": f"Short statement of {id}",
        "validation_status": status,
        "evidence_found": evidence,
        "notes": notes,
    }


def default_sections():
    """Two curated sections, every status represented at least once."""
    return {
        "Architecture, Compute & Infrast": [
            item("ARCH-001", "Needs Review"),
            item("ARCH-002", "Fail"),
            item("ARCH-003", "Pass", priority="nice_to_have"),
        ],
        "Enterprise-grade Security, Perm": [
            item("SEC-001", "Partial"),
            item("SEC-002", "Not Applicable", priority="nice_to_have"),
            item("SEC-003", "Pass"),
        ],
    }


_AUTO = object()


def build_checklist(
    path,
    sections=None,
    *,
    metadata=_AUTO,
    overall=_AUTO,
    include_total=True,
    total_override=None,
    breakdown=_AUTO,
    critical=_AUTO,
    critical_cols=("ID", "Section", "Title", "Status"),
    other=_AUTO,
    other_cols=("ID", "Section", "Title", "Status"),
    recommendations=("1. Fix ARCH-002 first.", "2. Review ARCH-001 with the customer."),
    recommendations_table=None,
    item_columns=None,
):
    """Write a checklist .xlsx to `path` and return the path.

    Each Summary block defaults to values derived from `sections` (so a default
    build is fully self-consistent and yields no data warnings); pass an
    explicit value to override a block, or None to omit that block entirely.
    `critical`/`other` take lists of row dicts keyed by lowercased column name.
    """
    sections = default_sections() if sections is None else sections
    all_items = [(sheet, it) for sheet, items in sections.items() for it in items]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value="Dataiku DSS Diagnosis - Checklist Review").font = Font(bold=True, size=14)

    if metadata is _AUTO:
        metadata = {
            "Bundle:": "synthetic_bundle_2026-07-22",
            "Node / Version:": "Design node / 14.4.3",
            "Report generated:": "2026-07-22 10:00",
        }
    row = 3
    for label, value in (metadata or {}).items():
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=11)
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    def block_header(text):
        nonlocal row
        ws.cell(row=row, column=1, value=text).font = BLOCK_FONT
        row += 1

    def table_header(values):
        nonlocal row
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = TABLE_HEADER_FONT
            cell.fill = TABLE_HEADER_FILL
        row += 1

    def write_row(values):
        nonlocal row
        for c, v in enumerate(values, start=1):
            ws.cell(row=row, column=c, value=v)
        row += 1

    # Overall Status Counts is the anchor every other block header is matched
    # against by style, so it's always written.
    if overall is _AUTO:
        overall = {s: sum(1 for _, it in all_items if it["validation_status"] == s) for s in STATUSES}
    block_header("Overall Status Counts")
    table_header(["Status", "All items"])
    for status, count in (overall or {}).items():
        write_row([status, count])
    if include_total:
        write_row(["Total", total_override if total_override is not None else len(all_items)])
    row += 1

    if breakdown is _AUTO:
        breakdown = {
            sheet.strip(): {s: sum(1 for it in items if it["validation_status"] == s) for s in STATUSES}
            for sheet, items in sections.items()
        }
    if breakdown is not None:
        block_header("Per-Section Breakdown")
        table_header(["Section"] + STATUSES)
        for section, counts in breakdown.items():
            write_row([section] + [counts.get(s, 0) for s in STATUSES])
        row += 1

    def id_led_rows(status_filter):
        return [
            {"id": it["id"], "section": sheet, "title": it["title"], "status": it["validation_status"]}
            for sheet, it in all_items
            if it["priority"] == "must_have" and it["validation_status"] in status_filter
        ]

    if critical is _AUTO:
        critical = id_led_rows({"Fail"})
    if critical is not None:
        block_header("Critical Findings - Must-Have Items Failing")
        table_header(list(critical_cols))
        for r in critical:
            write_row([r.get(c.lower(), "") for c in critical_cols])
        row += 1

    if other is _AUTO:
        other = id_led_rows({"Partial", "Needs Review"})
    if other is not None:
        block_header("Other Must-Have Items: Partial / Needs Review")
        table_header(list(other_cols))
        for r in other:
            write_row([r.get(c.lower(), "") for c in other_cols])
        row += 1

    if recommendations_table is not None:
        block_header("Priority-Ordered Recommendations")
        for r in recommendations_table:
            write_row(list(r))
    elif recommendations is not None:
        block_header("Priority-Ordered Recommendations")
        for rec in recommendations:
            write_row([rec])

    columns = item_columns or read_checklist.REQUIRED_ITEM_COLUMNS
    for sheet, items in sections.items():
        sws = wb.create_sheet(sheet)
        sws.append(list(columns))
        for it in items:
            sws.append([it.get(col, f"{col} value") for col in columns])

    wb.save(path)
    return Path(path)


@pytest.fixture
def checklist_factory(tmp_path):
    counter = {"n": 0}

    def make(*args, **kwargs):
        counter["n"] += 1
        return build_checklist(tmp_path / f"checklist_{counter['n']}.xlsx", *args, **kwargs)

    return make


@pytest.fixture
def section_config():
    import common

    return common.load_config("section_names.yaml")


def pytest_addoption(parser):
    parser.addoption(
        "--update-golden", action="store_true", default=False,
        help="Rewrite tests/golden/*.json from the current output instead of comparing against it.",
    )
