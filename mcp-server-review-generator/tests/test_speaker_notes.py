"""Every registered style must carry each flagged item's full Evidence/Notes in speaker notes.

Parametrized over `styles.STYLES`, so a new style fails here until it attaches notes
(use `deck_shared.finding_note_lines()`)."""

import pytest

import common
import styles
from conftest import item

pytest.importorskip("pptx")

LONG_EVIDENCE = "Evidence for {id}: " + "a very long observation that no table cell shows in full; " * 12 + "END-{id}"


def _sections():
    flagged = {"ARCH-001": "Needs Review", "ARCH-002": "Fail", "SEC-001": "Partial"}
    mk = lambda i: item(i, flagged[i], evidence=LONG_EVIDENCE.format(id=i), notes=f"Notes for {i}\n- point of {i}")
    return {
        "Architecture, Compute & Infrast": [mk("ARCH-001"), mk("ARCH-002"), item("ARCH-003", "Pass", priority="nice_to_have")],
        "Enterprise-grade Security, Perm": [mk("SEC-001"), item("SEC-003", "Pass")],
    }, flagged


@pytest.mark.skipif(not common.BRANDING_TEMPLATE.exists(), reason="branding template not present")
@pytest.mark.parametrize("style", sorted(styles.STYLES))
def test_style_attaches_full_evidence_and_notes_to_speaker_notes(style, checklist_factory, tmp_path):
    from pptx import Presentation

    sections, flagged = _sections()
    out = tmp_path / f"{style}.pptx"
    styles.build(style, styles.BuildRequest(checklist_factory(sections), "Acme", out, common.BRANDING_TEMPLATE))
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(str(out)).slides if s.has_notes_slide]
    for item_id in flagged:
        holding = [n for n in notes if f"{item_id}:" in n]
        assert holding, f"{style}: no speaker notes mention {item_id}"
        assert any(f"END-{item_id}" in n and f"point of {item_id}" in n for n in holding), \
            f"{style}: notes for {item_id} lack its full evidence/notes text"
