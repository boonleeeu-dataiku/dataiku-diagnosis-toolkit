"""Write the Summary sheet of a completed checklist deterministically.

The mechanical parts of the Summary (metadata labels, the five block headers
and their one shared style, status counts, per-section tallies, the
must-have finding rows and their ID/Section/Title/Status columns) are computed
here from the section sheets, so the layout read_checklist.py expects is
identical on every run, whichever LLM drove the review. The caller supplies
only the judgment text: a one-line key point per listed finding and the
ordered recommendations.

The block headers are recognised by read_checklist.find_section_blocks() via a
shared (bold, size, fill) signature, so BLOCK_HEADER_FONT/FILL below must not
be reused for any other column-A cell on the sheet.
"""

import logging
import os
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

import common
import read_checklist

logger = logging.getLogger(__name__)

STATUSES = ["Pass", "Fail", "Partial", "Needs Review", "Not Applicable"]
KEY_POINT_MAX_CHARS = 90
FINDING_COLUMNS = ["ID", "Section", "Title", "Status", "Key point"]

HEADERS = {
    "overall": "Overall Status Counts",
    "per_section": "Per-Section Breakdown",
    "critical": "Critical Findings - Must-Have Items Failing",
    "other": "Other Must-Have Items: Partial / Needs Review",
    "recommendations": "Priority-Ordered Recommendations",
}

# The block-header style. No other column-A cell may share this signature.
BLOCK_HEADER_FONT = Font(bold=True, size=12)
BLOCK_HEADER_FILL = PatternFill("solid", fgColor="FFD9E2F3")
TITLE_FONT = Font(bold=True, size=14)
LABEL_FONT = Font(bold=True, size=11)
TABLE_HEADER_FONT = Font(bold=True, size=11)
TABLE_HEADER_FILL = PatternFill("solid", fgColor="FFF2F2F2")
STATUS_FILLS = {
    "Pass": "FFC6EFCE",
    "Fail": "FFFFC7CE",
    "Partial": "FFFFEB9C",
    "Needs Review": "FFBDD7EE",
    "Not Applicable": "FFD9D9D9",
}
COLUMN_WIDTHS = {"A": 46, "B": 20, "C": 44, "D": 16, "E": 90, "F": 16}

_STATUS_BY_LOWER = {s.lower(): s for s in STATUSES}


class SummaryInputError(ValueError):
    """The caller's input can't produce a valid Summary (message says how to fix it)."""


def _canonical_status(raw, item_id: str) -> str:
    status = _STATUS_BY_LOWER.get(str(raw or "").strip().lower())
    if status is None:
        raise SummaryInputError(
            f"Item {item_id} has validation_status {raw!r}; expected one of {STATUSES}. "
            "Fill in every item's status before writing the Summary."
        )
    return status


def _load_items(wb) -> dict[str, list]:
    tab_names = [n for n in wb.sheetnames if n != "Summary"]
    if not tab_names:
        raise read_checklist.ChecklistFormatError("The workbook has no checklist-item sheets besides 'Summary'.")
    items_by_sheet = {name: read_checklist.parse_section_sheet(wb[name], name) for name in tab_names}
    seen = {}
    for sheet, items in items_by_sheet.items():
        for it in items:
            if it.id in seen:
                raise SummaryInputError(f"Duplicate item id {it.id!r} (in {seen[it.id]!r} and {sheet!r}).")
            seen[it.id] = sheet
    return items_by_sheet


def _tally(items) -> dict[str, int]:
    counts = {s: 0 for s in STATUSES}
    for it in items:
        counts[_canonical_status(it.validation_status, it.id)] += 1
    return counts


def _finding_rows(items_by_sheet, statuses: set, key_points: dict) -> list[list]:
    """Must-have items whose status is in `statuses`, ordered by `key_points`
    (caller's order: most causally central first)."""
    candidates = {}
    for sheet, items in items_by_sheet.items():
        for it in items:
            if it.is_must_have and _canonical_status(it.validation_status, it.id) in statuses:
                candidates[it.id] = (sheet, it)
    ordered_ids = [i for i in key_points if i in candidates]
    ordered_ids += [i for i in candidates if i not in key_points]  # surfaced as an error by the caller
    rows = []
    for item_id in ordered_ids:
        sheet, it = candidates[item_id]
        rows.append([it.id, sheet, it.title, _canonical_status(it.validation_status, it.id), key_points.get(it.id, "")])
    return rows


def _validate_inputs(items_by_sheet, key_points, recommendations, reviewer, bundle):
    problems = []
    if not str(reviewer or "").strip():
        problems.append("reviewer is required.")
    if not str(bundle or "").strip():
        problems.append("bundle is required.")
    if not recommendations or not any(str(r).strip() for r in recommendations):
        problems.append("recommendations must contain at least one action.")

    known = {it.id for items in items_by_sheet.values() for it in items}
    unknown = [i for i in key_points if i not in known]
    if unknown:
        problems.append(f"key_points has unknown item id(s): {unknown}.")

    listed = set()
    for statuses in ({"Fail"}, {"Partial", "Needs Review"}):
        for row in _finding_rows(items_by_sheet, statuses, key_points):
            listed.add(row[0])
    missing = [i for i in listed if not str(key_points.get(i, "")).strip()]
    if missing:
        problems.append(f"key_points is missing a line for must-have item(s): {sorted(missing)}.")
    not_listed = [i for i in key_points if i in known and i not in listed]
    if not_listed:
        problems.append(
            f"key_points lists item(s) that are not must-have Fail/Partial/Needs Review: {sorted(not_listed)}. "
            "Remove them."
        )
    too_long = [i for i, v in key_points.items() if len(str(v).strip()) > KEY_POINT_MAX_CHARS]
    if too_long:
        problems.append(f"key point for {sorted(too_long)} exceeds {KEY_POINT_MAX_CHARS} characters; shorten it.")
    multiline = [i for i, v in key_points.items() if "\n" in str(v)]
    if multiline:
        problems.append(f"key point for {sorted(multiline)} must be a single line.")
    if problems:
        raise SummaryInputError("Cannot write the Summary:\n- " + "\n- ".join(problems))


def _write_sheet(ws, *, items_by_sheet, key_points, recommendations, metadata):
    row = 1

    def put(values, font=None, fill=None, status_col=None):
        nonlocal row
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if status_col == c and v in STATUS_FILLS:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILLS[v])
        row += 1

    def block_header(text):
        nonlocal row
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = BLOCK_HEADER_FONT
        cell.fill = BLOCK_HEADER_FILL
        row += 1

    def blank():
        nonlocal row
        row += 1

    put(["Dataiku DSS Diagnosis - Checklist Review"], font=TITLE_FONT)
    blank()
    for label, value in metadata:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1
    blank()

    all_items = [it for items in items_by_sheet.values() for it in items]
    overall = _tally(all_items)
    block_header(HEADERS["overall"])
    for status in STATUSES:
        put([status, overall[status]])
    ws.cell(row=row, column=1, value="Total").font = LABEL_FONT
    ws.cell(row=row, column=2, value=len(all_items)).font = LABEL_FONT
    row += 1
    blank()

    block_header(HEADERS["per_section"])
    put(["Section"] + STATUSES, font=TABLE_HEADER_FONT, fill=TABLE_HEADER_FILL)
    for sheet, items in items_by_sheet.items():
        counts = _tally(items)
        put([sheet] + [counts[s] for s in STATUSES])
    blank()

    for key, statuses in (("critical", {"Fail"}), ("other", {"Partial", "Needs Review"})):
        block_header(HEADERS[key])
        put(FINDING_COLUMNS, font=TABLE_HEADER_FONT, fill=TABLE_HEADER_FILL)
        for finding in _finding_rows(items_by_sheet, statuses, key_points):
            finding[4] = str(finding[4]).strip()
            put(finding, status_col=4)
        blank()

    block_header(HEADERS["recommendations"])
    n = 0
    for rec in recommendations:
        text = common.strip_leading_number(str(rec).strip())
        if not text:
            continue
        n += 1
        put([f"{n}. {text}"])

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def _verify(path: Path, items_by_sheet, config, key_points, recommendations) -> None:
    """Re-read the saved file the way the deck generator will and compare it to
    what was meant to be written."""
    data = read_checklist.read_checklist(path, config)
    all_items = [it for items in items_by_sheet.values() for it in items]
    problems = []

    expected_counts = {**_tally(all_items), "Total": len(all_items)}
    if data.overall_counts != expected_counts:
        problems.append(f"overall counts read back as {data.overall_counts}, expected {expected_counts}.")
    # The reader strips cell text, and a tab name can carry trailing whitespace
    # (the bundled template's "Advanced Security Options (DSS ").
    by_stripped = {name.strip(): items for name, items in items_by_sheet.items()}
    if [b["section"] for b in data.per_section_breakdown] != list(by_stripped):
        problems.append("Per-Section Breakdown rows do not match the section tabs.")
    for block in data.per_section_breakdown:
        if block["counts"] != _tally(by_stripped.get(block["section"], [])):
            problems.append(f"Per-Section counts for {block['section']!r} are wrong.")
    for name, got, statuses in (
        ("Critical Findings", data.critical_findings, {"Fail"}),
        ("Other Must-Have Items", data.other_must_have, {"Partial", "Needs Review"}),
    ):
        expected_ids = [r[0] for r in _finding_rows(items_by_sheet, statuses, key_points)]
        if [r.get("id") for r in got] != expected_ids:
            problems.append(f"{name} read back as {[r.get('id') for r in got]}, expected {expected_ids}.")
    expected_recs = len([r for r in recommendations if common.strip_leading_number(str(r).strip())])
    if len(data.recommendations) != expected_recs:
        problems.append(f"{len(data.recommendations)} recommendations read back, expected {expected_recs}.")
    for label in ("Bundle", "Report generated", "Reviewer"):
        if label not in data.bundle_metadata:
            problems.append(f"metadata row {label!r} not read back.")
    if problems:
        raise RuntimeError("Summary round-trip check failed:\n- " + "\n- ".join(problems))


def write_summary(
    checklist_path,
    *,
    reviewer: str,
    bundle: str,
    node_version: str = "",
    diagnosis_generated: str = "",
    key_points: dict | None = None,
    recommendations: list | None = None,
    config: dict,
    today: date | None = None,
) -> dict:
    """Recreate the Summary sheet (first tab) of the checklist at `checklist_path`.

    key_points: {item id: one-line key point (<= 90 chars)} for every must-have
        Fail / Partial / Needs Review item. Dict order sets the order within each
        block (most causally central first).
    recommendations: ordered action strings, root cause before symptoms; numbering is
        added (or normalised) here.
    Returns counts and the block sizes written. Raises SummaryInputError for bad input
    (before touching the file) and RuntimeError if the written file fails the round-trip
    check (the original file is then left unchanged).
    """
    path = Path(checklist_path)
    key_points = {str(k).strip(): str(v).strip() for k, v in (key_points or {}).items()}
    recommendations = list(recommendations or [])

    wb = openpyxl.load_workbook(path)
    items_by_sheet = _load_items(wb)
    _validate_inputs(items_by_sheet, key_points, recommendations, reviewer, bundle)

    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    ws = wb.create_sheet("Summary", 0)
    wb.active = 0
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet is ws

    metadata = [
        ("Bundle:", str(bundle).strip()),
        ("Node / Version:", str(node_version or "").strip()),
        ("Diagnosis generated:", str(diagnosis_generated or "").strip()),
        ("Report generated:", (today or date.today()).isoformat()),
        ("Reviewer:", str(reviewer).strip()),
    ]
    _write_sheet(ws, items_by_sheet=items_by_sheet, key_points=key_points,
                 recommendations=recommendations, metadata=metadata)

    tmp = path.with_name(path.stem + ".summary-tmp" + path.suffix)
    wb.save(tmp)
    try:
        _verify(tmp, items_by_sheet, config, key_points, recommendations)
        os.replace(tmp, path)
    finally:
        if tmp.exists():
            tmp.unlink()

    all_items = [it for items in items_by_sheet.values() for it in items]
    return {
        "checklist_path": str(path),
        "counts": {**_tally(all_items), "Total": len(all_items)},
        "sections": list(items_by_sheet),
        "critical_findings": len(_finding_rows(items_by_sheet, {"Fail"}, key_points)),
        "other_must_have": len(_finding_rows(items_by_sheet, {"Partial", "Needs Review"}, key_points)),
        "recommendations": len([r for r in recommendations if common.strip_leading_number(str(r).strip())]),
    }
