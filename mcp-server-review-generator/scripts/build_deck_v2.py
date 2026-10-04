#!/usr/bin/env python3
"""v2 deck: verdict first, 3 risks, quick wins, asks, roadmap, then an appendix.

Pipeline: build the v1 deck into a temp file (it supplies the branded cover and
end card), then use python-pptx to add the v2 slides built from that deck's own
layouts, and drop every other v1 slide. python-pptx is used ONLY here -- the v1
path stays raw OOXML (see INTENT.md / CLAUDE.md).

Content comes from deck_analysis (rules) and an optional narrative.json
(human judgment, validated in narrative.py). Without a narrative, every slide
falls back to text derived from checklist cells; nothing is invented.
"""

import hashlib
import logging
import math
import re
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import build_deck
import common
import deck_analysis as da
import narrative as narr_mod
import read_checklist
import section_names

logger = logging.getLogger(__name__)

DK = "06312E"; MINT = "3EDAB2"; BG = "FEFEF9"; INK = "1A1A1A"; GREY = "5B6470"
LINE = "E2E5EA"; RED = "C0392B"; AMBER = "EDAB4F"; BLUE = "7092F2"; WHITE = "FFFFFF"
PALE = "D6E4E2"
STATUS = {  # chip fill, text
    "Pass": ("D9F7EF", "0B6B55"), "Fail": ("FDEAEA", "9B2C2C"), "Partial": ("FFF1D6", "8A5A00"),
    "Needs Review": ("E8EEFB", "2F4C9E"), "Not Applicable": ("F1F2F4", "5B6470"),
}
SEG = {"Pass": MINT, "Needs Review": BLUE, "Partial": AMBER, "Fail": RED}
DOT = {"ok": MINT, "watch": AMBER, "neutral": "B8BFC9"}
L_CREAM, L_DARK = "CUSTOM_5", "CUSTOM_3_1"


def RGB(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- helpers
class Deck:
    def __init__(self, prs):
        self.prs = prs
        self.layouts = {l.name: l for l in prs.slide_layouts}

    def cream(self, title, size=20):
        s = self.prs.slides.add_slide(self.layouts[L_CREAM])
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = RGB(BG)
        set_title(s, title, size)
        return s

    def dark(self, title, size=22, w=9.22):
        s = self.prs.slides.add_slide(self.layouts[L_DARK])
        set_title(s, title, size, w=w)
        return s


def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def card(slide, x, y, w, h, fill=WHITE, line=LINE, r=0.05, name=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(sh); sh.adjustments[0] = r
    sh.fill.solid(); sh.fill.fore_color.rgb = RGB(fill)
    if line:
        sh.line.color.rgb = RGB(line); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    if name:
        sh.name = name
    return sh


def rect(slide, x, y, w, h, fill, name=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(sh); sh.fill.solid(); sh.fill.fore_color.rgb = RGB(fill); sh.line.fill.background()
    if name:
        sh.name = name
    return sh


def dot(slide, x, y, d, color, name):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    strip_style(sh); sh.fill.solid(); sh.fill.fore_color.rgb = RGB(color); sh.line.fill.background()
    sh.name = name
    return sh


def badge(slide, x, y, d, text, name, size=9):
    sh = dot(slide, x, y, d, DK, name)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(tf, [P(text, size, True, WHITE, "c")])
    return sh


def P(runs, size=11, bold=False, color=INK, align="l", after=0, font="Roboto", before=0):
    if isinstance(runs, str):
        runs = [(runs, {})]
    return dict(runs=runs, size=size, bold=bold, color=color, align=align, after=after, font=font, before=before)


def fill_tf(tf, paras):
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[p["align"]]
        para.space_after = Pt(p["after"]); para.space_before = Pt(p["before"])
        for text, o in p["runs"]:
            r = para.add_run(); r.text = text
            f = r.font; f.name = o.get("font", p["font"]); f.size = Pt(o.get("size", p["size"]))
            f.bold = o.get("bold", p["bold"]); f.color.rgb = RGB(o.get("color", p["color"]))


def tb(slide, x, y, w, h, paras, anchor="t", name=None):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        t.name = name
    tf = t.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras)
    return t


def set_title(slide, text, size=20, w=9.22):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.39), Inches(0.28), Inches(w), Inches(0.75)
    tf = t.text_frame; tf.word_wrap = True
    tf.text = text
    for para in tf.paragraphs:
        for r in para.runs:
            r.font.size = Pt(size)


def chip(slide, x, y, w, h, text, status=None, fill=None, color=None, size=8, name=None):
    f, c = STATUS[status] if status else (fill, color)
    sh = card(slide, x, y, w, h, fill=f, line=None, r=0.3, name=name)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(tf, [P(text, size, True, c, "c")])
    return sh


def cell_set(cell, paras, fill=None):
    cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.05)
    cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame; tf.word_wrap = True
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    fill_tf(tf, paras)
    if fill:
        cell.fill.solid(); cell.fill.fore_color.rgb = RGB(fill)
    else:
        cell.fill.background()
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    for tag, vis in (("a:lnL", False), ("a:lnR", False), ("a:lnT", False), ("a:lnB", True)):
        ln = etree.SubElement(tcPr, qn(tag), w="9525" if vis else "0")
        if vis:
            sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=LINE)
        else:
            etree.SubElement(ln, qn("a:noFill"))
    # schema order: borders must precede fill
    for e in [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]:
        tcPr.remove(e); tcPr.append(e)


def table(slide, x, y, widths, header, body, row_h, hdr_h=0.3, name="Table"):
    """row_h: a number, or one height per body row."""
    heights = row_h if isinstance(row_h, list) else [row_h] * len(body)
    gs = slide.shapes.add_table(len(body) + 1, len(widths), Inches(x), Inches(y),
                                Inches(sum(widths)), Inches(hdr_h + sum(heights)))
    gs.name = name
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(hdr_h)
    for j, h in enumerate(header):
        cell_set(tbl.cell(0, j), [P(h, 8.5, True, DK)], fill=MINT)
    for i, r in enumerate(body, start=1):
        tbl.rows[i].height = Inches(heights[i - 1])
        for j, c in enumerate(r):
            cell_set(tbl.cell(i, j), c if isinstance(c, list) else [P(c, 9)], fill=WHITE)
    return tbl


def clip(text, n):
    text = (text or "").strip()
    return text if len(text) <= n else text[:n - 3].rsplit(" ", 1)[0] + "..."


def est_lines(text, width_in, size_pt):
    chars = max(1, int(width_in * 72 / (size_pt * 0.55)))
    return sum(max(1, math.ceil(len(l) / chars)) for l in (text or " ").split("\n"))


# ------------------------------------------------------------ slide content
def verdict_slide(d, a, narr):
    c = a.counts
    f, nr = c["Fail"], c["Needs Review"]
    title = narr.get("verdict_title") or (
        f"{c['Pass']} of {a.applicable} applicable checks pass; {f} fail and {nr} need confirmation")
    s = d.dark(title, 22)
    card(s, 0.39, 1.4, 2.55, 1.55, name="Headline stat card")
    tb(s, 0.39, 1.5, 2.55, 0.75, [P(f"{c['Pass']} of {a.applicable}", 34, True, DK, "c", font="Spectral")], "m", name="Headline stat")
    tb(s, 0.55, 2.3, 2.25, 0.55, [P("applicable checks pass", 10.5, False, INK, "c"),
                                  P(f"{a.na} of {a.total} items are N/A and excluded", 8.5, False, GREY, "c")], name="Headline caption")
    bx, bw, by, bh = 3.2, 6.4, 1.55, 0.5
    x = bx
    for label in da.STATUS_ORDER:
        n = c[label]
        if not n:
            continue
        w = bw * n / max(a.applicable, 1)
        rect(s, x, by, w, bh, SEG[label], name=f"Bar - {label}")
        tb(s, x, by, w, bh, [P(str(n), 14, True, DK if label in ("Pass", "Partial") else WHITE, "c")], "m", name=f"Bar count - {label}")
        cw = max(w, 0.9)
        tb(s, x + w / 2 - cw / 2, by + bh + 0.06, cw, 0.2, [P(label, 9, False, PALE, "c")], name=f"Bar label - {label}")
        x += w
    hl = narr.get("highlights")
    if not hl:
        parts = []
        for disp, _ in a.section_order:
            sc = a.section_counts(disp)
            if sc["Fail"]:
                parts.append(f"{disp}: {sc['Fail']} fail")
        hl = "  ·  ".join(parts) if parts else "No section has a failing check"
    tb(s, bx, 2.45, bw, 0.45, [P(hl, 9.5, False, PALE)], name="Section highlights")

    tk = narr.get("takeaways") or fallback_takeaways(a)
    cols = {"good": MINT, "risk": RED, "win": AMBER}
    for i, t in enumerate(tk):
        cx = 0.39 + i * 3.1
        card(s, cx, 3.2, 3.0, 1.6, name=f"Takeaway card {i + 1}")
        dot(s, cx + 0.18, 3.38, 0.16, cols[t.get("tone", "good")], f"Takeaway dot {i + 1}")
        tb(s, cx + 0.42, 3.33, 2.45, 0.3, [P(t["heading"], 12, True, DK)], "m", name=f"Takeaway head {i + 1}")
        tb(s, cx + 0.18, 3.78, 2.65, 0.95, [P(t["body"], 10.5)], name=f"Takeaway body {i + 1}")


def fallback_takeaways(a):
    c = a.counts
    passed = a.by_status("Pass")
    sound = ", ".join(clip(r.title, 28) for r in passed[:3]) or "No passing checks"
    risk_ids = ", ".join(a.root_causes[0][:4]) if a.root_causes else ", ".join(r.id for r in a.by_status("Fail")[:4])
    return [
        {"heading": "What is sound", "tone": "good", "body": f"{c['Pass']} checks pass, including {sound}."},
        {"heading": "What is at risk", "tone": "risk",
         "body": f"{c['Fail']} fail and {c['Partial']} are partial." + (f" Largest cluster: {risk_ids}." if risk_ids else "")},
        {"heading": "Quick wins", "tone": "win",
         "body": f"{len(a.quick_wins)} of the {c['Fail']} fails are single-setting changes." if c["Fail"] else "No failing checks."},
    ]


def snapshot_slide(d, a, narr, data):
    cards = narr.get("snapshot")
    if not cards:
        md = data.bundle_metadata
        cards = []
        for label, key in (("NODE / VERSION", "Node / Version"), ("BUNDLE", "Bundle"), ("DIAGNOSIS GENERATED", "Diagnosis generated")):
            if md.get(key):
                cards.append({"label": label, "value": clip(str(md[key]), 28), "note": "", "state": "neutral"})
        for disp, _ in a.section_order:
            sc = a.section_counts(disp)
            app = sum(sc[k] for k in da.STATUS_ORDER)
            cards.append({"label": clip(disp, 26).upper(), "value": f"{sc['Pass']} of {app} pass",
                          "note": f"{sc['Fail']} fail, {sc['Partial']} partial", "state": "watch" if sc["Fail"] else "ok"})
        cards = cards[:8]
    title = narr.get("snapshot_title") or "Instance snapshot"
    s = d.cream(title)
    for i, c in enumerate(cards):
        cx = 0.39 + (i % 4) * 2.34; cy = 1.3 + (i // 4) * 1.75
        card(s, cx, cy, 2.2, 1.6, name=f"Snapshot card {i + 1}")
        tb(s, cx + 0.16, cy + 0.16, 1.7, 0.2, [P(c["label"], 8, True, GREY)], name=f"Snapshot label {i + 1}")
        dot(s, cx + 1.92, cy + 0.18, 0.14, DOT[c.get("state", "neutral")], f"Snapshot status {i + 1}")
        tb(s, cx + 0.16, cy + 0.55, 1.9, 0.45, [P(c["value"], 15, True, DK, font="Spectral")], "m", name=f"Snapshot value {i + 1}")
        tb(s, cx + 0.16, cy + 1.08, 1.9, 0.4, [P(c.get("note", ""), 9.5, False, GREY)], name=f"Snapshot note {i + 1}")
    for k, (col, txt) in enumerate([(MINT, "in line with guidance"), (AMBER, "worth attention"), ("B8BFC9", "context")]):
        lx = 0.39 + k * 2.1
        dot(s, lx, 4.93, 0.12, col, f"Legend dot {k + 1}")
        tb(s, lx + 0.2, 4.9, 1.7, 0.2, [P(txt, 8.5, False, GREY)], name=f"Legend text {k + 1}")


def risk_title(n, title):
    """Slide title for risk card `n`: a narrative title without the "Risk N:" lead gets one."""
    title = str(title).strip()
    return title if re.match(rf"Risk\s*{n}\s*:", title, re.I) else f"Risk {n}: {title}"


def risk1_slide(d, a, narr):
    r1 = narr.get("risk1")
    if r1 is None:
        r1 = fallback_risk1(a)
    if not r1:
        return None
    s = d.cream(risk_title(1, r1["title"]))
    tiles = r1.get("tiles", [])[:5]
    for i, t in enumerate(tiles):
        cx = 0.39 + (i % 2) * 2.45; cy = 1.2 + (i // 2) * 1.2
        card(s, cx, cy, 2.35, 1.1, name=f"Evidence tile {i + 1}")
        tb(s, cx + 0.15, cy + 0.08, 0.95, 0.55, [P(str(t["number"]), 24, True, RED, font="Spectral")], "m", name=f"Evidence number {i + 1}")
        tb(s, cx + 1.05, cy + 0.1, 1.2, 0.5, [P(t["label"], 9.5, True, DK)], "m", name=f"Evidence label {i + 1}")
        note = f"{t['id']}: {t['note']}" if t.get("id") and t.get("note") else (t.get("id") or t.get("note", ""))
        tb(s, cx + 0.15, cy + 0.68, 2.1, 0.4, [P(note, 8.5, False, GREY)], name=f"Evidence note {i + 1}")
    ids = sorted({i for st in r1.get("fix", []) for i in (st["id"] if isinstance(st["id"], list) else [st["id"]])})
    if r1.get("linked_label") or ids:
        card(s, 0.39 + 2.45, 1.2 + 2 * 1.2, 2.35, 1.1, fill=DK, line=None, name="Root cause card")
        tb(s, 0.39 + 2.45 + 0.15, 1.2 + 2 * 1.2 + 0.1, 2.05, 0.9,
           [P("One root cause", 10, True, MINT, after=3),
            P(r1.get("linked_label") or f"{len(ids)} checks, one sequence of fixes.", 10, False, WHITE)], "m", name="Root cause text")
    card(s, 5.45, 1.2, 4.15, 1.05, fill="FDEAEA", line="E8A0A0", name="Risk card")
    tb(s, 5.62, 1.28, 3.85, 0.9, [P("Risk", 9, True, "9B2C2C", after=2), P(r1["risk"], 11)], "m", name="Risk text")
    card(s, 5.45, 2.4, 4.15, 2.4, name="Fix card")
    tb(s, 5.62, 2.5, 3.85, 0.25, [P("Fix, in this order", 11, True, DK)], name="Fix heading")
    fix = r1.get("fix", [])[:5]
    step = min(0.4, 1.6 / max(len(fix), 1))
    for i, st in enumerate(fix, start=1):
        fy = 2.88 + (i - 1) * step
        badge(s, 5.62, fy, 0.27, str(i), f"Fix step badge {i}")
        tb(s, 6.02, fy, 2.7, 0.27, [P(st["text"], 10)], "m", name=f"Fix step {i}")
        sid = ", ".join(st["id"]) if isinstance(st["id"], list) else st["id"]
        tb(s, 8.8, fy, 0.7, 0.27, [P(sid, 8.5, True, GREY, "r")], "m", name=f"Fix item {i}")
    if r1.get("caveat"):
        tb(s, 5.62, 4.5, 3.85, 0.25, [P(r1["caveat"], 8.5, False, GREY)], name="Fix caveat")


def fallback_risk1(a):
    """Largest root-cause group (or, failing that, the Fail items) from cells only."""
    ids = a.root_causes[0] if a.root_causes else [r.id for r in a.by_status("Fail")][:5]
    if not ids:
        return None
    rows = sorted((a.rows[i] for i in ids), key=lambda r: (r.status != "Fail", r.priority != "Must", r.id))
    tiles = []
    for r in rows[:5]:
        m = next((re.search(r"\d[\d,.]*", t) for t in [r.headline] + r.evidence if re.search(r"\d", t)), None)
        tiles.append({"number": m.group(0) if m else r.status, "label": clip(r.title, 38), "id": r.id,
                      "note": clip(r.headline, 60)})
    return {"title": f"Risk 1: {len(rows)} linked checks share one root cause" if a.root_causes else "Risk 1: failing checks",
            "tiles": tiles, "risk": clip(rows[0].headline, 110),
            "fix": [{"text": clip(r.action or r.headline, 52), "id": r.id} for r in rows[:5]]}


def risk2_slide(d, a, narr):
    r2 = narr.get("risk2")
    if r2 is None:
        used = {i for g in a.root_causes[:1] for i in g}
        fails = [r for r in a.by_status("Fail") if r.id not in used]
        if not fails:
            return None
        sec = max({r.section for r in fails}, key=lambda s: sum(1 for r in fails if r.section == s))
        sel = [r for r in fails if r.section == sec]
        r2 = {"title": f"Risk 2: {len(sel)} failing checks in {sec}",
              "table": [{"id": r.id, "setting": clip(r.title, 38), "today": clip(r.headline, 50),
                         "target": clip(r.action or "-", 50)} for r in sel]}
    rows_ = r2["table"][:7]
    s = d.cream(risk_title(2, r2["title"]))
    body = []
    for e in rows_:
        row = a.rows[e["id"]]
        must = row.priority == "Must"
        body.append([[P(e["id"], 8.5, True, RED)], [P(e["setting"], 9.5, True)], [P(e["today"], 9.5)], [P(e["target"], 9.5)],
                     [P("Must have" if must else "Nice to have", 8.5, must, "9B2C2C" if must else GREY)]])
    table(s, 0.39, 1.2, [1.0, 2.2, 2.3, 2.5, 1.22], ["Item", "Setting", "Today", "Target", "Priority"], body, 0.41, name="Hardening table")
    pos = r2.get("positives")
    if pos is None:
        passed = [r for r in a.by_status("Pass") if r.section == a.rows[rows_[0]["id"]].section][:4]
        pos = ", ".join(clip(r.headline or r.title, 40) for r in passed)
    if pos:
        card(s, 0.39, 4.4, 9.21, 0.55, fill="E6F8F2", line=None, name="Positive note card")
        tb(s, 0.55, 4.4, 8.9, 0.55, [P([("Already in place: ", {"bold": True, "color": "0B6B55"}), (pos, {})], 10)], "m", name="Positive note")


def id_list(value):
    """A narrative `ids` field is a string or a list; always return a list of IDs."""
    return da.ID_RE.findall(value) if isinstance(value, str) else list(value)


def risk3_slide(d, a, narr):
    r3 = narr.get("risk3")
    if r3 is None:
        used = {i for g in a.root_causes[:1] for i in g}
        left = [r for r in a.rows.values() if r.status in ("Partial", "Needs Review") and r.id not in used]
        left.sort(key=lambda r: (r.status != "Partial", r.priority != "Must", r.id))
        if not left:
            return None
        r3 = {"title": "Risk 3: open items still to close",
              "cards": [{"heading": clip(r.title, 32), "ids": [r.id], "found": r.headline, "todo": r.action or "Confirm with the owner."}
                        for r in left[:3]]}
    s = d.cream(risk_title(3, r3["title"]))
    cards = r3["cards"][:3]
    for i, c in enumerate(cards):
        cx = 0.39 + i * 3.1
        c = {**c, "ids": id_list(c["ids"])}
        first = a.rows[c["ids"][0]]
        st = first.status
        label = "Pass, with caveat" if st == "Pass" else st
        card(s, cx, 1.25, 3.0, 3.2, name=f"Resilience card {i + 1}")
        tb(s, cx + 0.18, 1.4, 2.65, 0.3, [P(c["heading"], 12.5, True, DK)], "m", name=f"Resilience heading {i + 1}")
        tb(s, cx + 0.18, 1.78, 1.6, 0.2, [P(" / ".join(c["ids"]), 8.5, True, GREY)], name=f"Resilience ids {i + 1}")
        chip(s, cx + 1.72, 1.77, 1.1, 0.22, label, st, name=f"Resilience status {i + 1}")
        tb(s, cx + 0.18, 2.2, 2.65, 0.2, [P("WHAT WE FOUND", 8, True, GREY)], name=f"Found label {i + 1}")
        tb(s, cx + 0.18, 2.42, 2.65, 0.9, [P(c["found"], 10.5)], name=f"Found text {i + 1}")
        tb(s, cx + 0.18, 3.5, 2.65, 0.2, [P("WHAT TO DO", 8, True, GREY)], name=f"Act label {i + 1}")
        tb(s, cx + 0.18, 3.72, 2.65, 0.9, [P(c["todo"], 10.5, True, DK)], name=f"Act text {i + 1}")


def quickwins_slide(d, a, narr):
    qn_ = narr.get("quick_wins")
    if qn_ is None:
        if not a.quick_wins:
            return None
        qn_ = {"rows": [{"id": q.id, "setting": q.setting, "from": q.from_, "to": q.to, "confirm": q.confirm} for q in a.quick_wins]}
    rows_ = qn_["rows"][:6]
    fails = a.counts["Fail"]
    after = fails - len(rows_)
    s = d.cream(qn_.get("title") or f"Quick wins: one change window closes {len(rows_)} of the {fails} fails")
    body = [[[P(q["id"], 8.5, True, RED)], [P(q["setting"], 8.5, False, INK, font="Courier New")],
             [P([(q["from"], {"color": "9B2C2C"}), ("  →  ", {"color": GREY}), (q["to"], {"color": "0B6B55", "bold": True})], 9.5)],
             [P(q.get("confirm", "None"), 9, False, GREY)]] for q in rows_]
    table(s, 0.39, 1.25, [1.0, 2.55, 1.75, 1.25], ["Item", "Setting", "Change", "Confirm first"], body, min(0.62, 3.0 / max(len(body), 1)), name="Quick wins table")
    card(s, 7.15, 1.25, 2.45, 3.4, fill=DK, line=None, name="Before-after card")
    tb(s, 7.3, 1.4, 2.15, 0.25, [P("FAILS AFTER THE WINDOW", 8, True, MINT, "c")], name="Before-after label")
    tb(s, 7.3, 1.85, 2.15, 0.9, [P([(str(fails), {"color": "F3A39A"}), ("  →  ", {"color": MINT, "size": 20}), (str(after), {"color": WHITE})],
                                    34, True, WHITE, "c", font="Spectral")], "m", name="Before-after numbers")
    won = {q["id"] for q in rows_}
    rest = [r.id for r in a.by_status("Fail") if r.id not in won]
    tb(s, 7.3, 3.0, 2.15, 1.5, [P("Remaining fails:", 9, True, MINT, after=3)] +
       [P(i, 9, False, WHITE, after=2) for i in rest[:8]], name="Remaining fails")


def owners_slide(d, a, narr):
    g = (narr.get("owners") or {}).get("groups")
    notes = (narr.get("owners") or {}).get("notes", {})
    if g is None:
        g = [{"owner": o, "asks": [{"id": i, "ask": clip(a.rows[i].action or a.rows[i].title, 70)} for i in ids]} for o, ids in a.owners]
    g = [x for x in g if x["asks"]]
    if not g:
        return None
    total = sum(len(x["asks"]) for x in g)
    if total != a.counts["Needs Review"]:
        raise RuntimeError("owner grouping must equal the Needs Review total")
    s = d.cream(narr.get("owners_title") or f"What we need from you: {total} items to confirm, grouped by owner")
    n = min(len(g), 5)
    gap = 0.13
    w = (9.22 - gap * (n - 1)) / n
    if len(g) > n:  # fold overflow owners into the last column
        g = g[:n - 1] + [{"owner": "Other", "asks": [e for x in g[n - 1:] for e in x["asks"]]}]
    for i, grp in enumerate(g):
        cx = 0.39 + i * (w + gap)
        card(s, cx, 1.2, w, 3.85, name=f"Owner card {i + 1}")
        tb(s, cx + 0.14, 1.3, w - 0.7, 0.3, [P(grp["owner"], 11, True, DK)], "m", name=f"Owner name {i + 1}")
        chip(s, cx + w - 0.46, 1.34, 0.32, 0.22, str(len(grp["asks"])), fill="E8EEFB", color="2F4C9E", size=9, name=f"Owner count {i + 1}")
        step = min(0.54, 2.7 / len(grp["asks"]))
        y = 1.72
        for ask in grp["asks"]:
            tb(s, cx + 0.14, y, w - 0.25, 0.2, [P(ask["id"], 8, True, "2F4C9E")], name=f"Owner item id {ask['id']}")
            tb(s, cx + 0.14, y + 0.17, w - 0.25, 0.3, [P(ask["ask"], 9)], name=f"Owner item ask {ask['id']}")
            y += step
        if notes.get(grp["owner"]):
            tb(s, cx + 0.14, 4.47, w - 0.25, 0.5, [P(notes[grp["owner"]], 8, False, GREY)], name=f"Owner note {i + 1}")


def roadmap_slide(d, a, narr, data):
    road = narr.get("roadmap")
    if road is None:
        recs = [da.LEADING_NUMBER_RE.sub("", r) for r in data.recommendations]
        if not recs:
            return None
        per = math.ceil(len(recs) / 3)
        cols = [recs[i * per:(i + 1) * per] for i in range(3)]
        road = {k: [{"effort": "", "action": clip(re.sub(r"\s*\([^)]*\)\s*\.?$", "", t).rstrip("."), 70), "ids": da.expand_ids(t)} for t in c]
                for k, c in zip(("now", "next", "plan"), cols)}
    s = d.cream(road.get("title") or "Roadmap: what to do now, next and later")
    spec = [("NOW", "This month", "now", RED), ("NEXT", "This quarter", "next", AMBER), ("PLAN", "Next two quarters", "plan", BLUE)]
    for i, (nm, when, key, col) in enumerate(spec):
        cx = 0.39 + i * 3.1
        items = road.get(key, [])[:5]
        card(s, cx, 1.2, 3.0, 3.55, name=f"Roadmap column {nm}")
        dot(s, cx + 0.16, 1.37, 0.16, col, f"Roadmap dot {nm}")
        tb(s, cx + 0.42, 1.3, 1.1, 0.3, [P(nm, 13, True, DK)], "m", name=f"Roadmap name {nm}")
        tb(s, cx + 1.3, 1.3, 1.55, 0.3, [P(when, 9, False, GREY, "r")], "m", name=f"Roadmap when {nm}")
        step = min(0.72, 2.8 / max(len(items), 1))
        y = 1.8
        for j, e in enumerate(items, start=1):
            eff = str(e.get("effort", "")).upper()
            if eff:
                chip(s, cx + 0.16, y + 0.03, 0.3, 0.3, eff, fill=DK, color=WHITE, size=9, name=f"Roadmap effort {nm}-{j}")
            tb(s, cx + 0.58, y, 2.3, 0.4, [P(e["action"], 9.5, True)], name=f"Roadmap action {nm}-{j}")
            ids = ", ".join(id_list(e["ids"]))
            tb(s, cx + 0.58, y + 0.4, 2.3, 0.2, [P(ids, 8, False, GREY)], name=f"Roadmap ids {nm}-{j}")
            y += step
    foot = road.get("footnote") or ("Effort is indicative: S about a day, M a few days, L weeks. Owners and dates to be agreed with the customer."
                                    if any(e.get("effort") for k in ("now", "next", "plan") for e in road.get(k, [])) else
                                    "Order follows the checklist's priority-ordered recommendations. Owners and dates to be agreed with the customer.")
    tb(s, 0.39, 4.85, 9.2, 0.4, [P(foot, 8.5, False, GREY)], name="Roadmap footnote")


def appendix_divider(d):
    s = d.dark("Appendix", 28, w=6)
    tb(s, 0.39, 1.7, 8, 1.2, [P(t, 12, False, PALE, after=3) for t in
                              ("How the review was done", "Passed checks", "Findings by section", "Not applicable items")], name="Appendix contents")


def method_slide(d, a, data):
    md = data.bundle_metadata
    s = d.cream("How the review was done")
    when = md.get("Diagnosis generated")
    steps = [("1", "Diagnostic bundle", f"Exported from the DSS instance{' on ' + str(when) if when else ''}"),
             ("2", "Standard checklist", f"{a.total} items across {len(a.section_order)} sections: "
              + ", ".join(x[0].split(",")[0].split(" ")[0] for x in a.section_order) + "."),
             ("3", "Status per item", "Each item is scored from the evidence in the bundle. No live checks were run.")]
    for i, (n, h, b) in enumerate(steps):
        cx = 0.39 + i * 3.1
        card(s, cx, 1.25, 2.9, 1.5, name=f"Method step {n}")
        badge(s, cx + 0.16, 1.4, 0.3, n, f"Method badge {n}", 10)
        tb(s, cx + 0.58, 1.4, 2.2, 0.3, [P(h, 11.5, True, DK)], "m", name=f"Method head {n}")
        tb(s, cx + 0.16, 1.9, 2.6, 0.8, [P(b, 9.5)], name=f"Method body {n}")
    defs = [("Pass", "Evidence shows the item is met."), ("Partial", "Met in part; some evidence shows a gap."),
            ("Fail", "Evidence shows the item is not met."), ("Needs Review", "Cannot be verified from the bundle; needs customer confirmation."),
            ("Not Applicable", "Not relevant to this deployment.")]
    for i, (st, text) in enumerate(defs):
        y = 3.05 + i * 0.36
        chip(s, 0.39, y, 1.25, 0.26, st, st, size=9, name=f"Status chip {st}")
        tb(s, 1.8, y, 7.8, 0.26, [P(text, 10)], "m", name=f"Status definition {st}")


def passed_slide(d, a, v2):
    s = d.cream(f"Passed checks: {a.counts['Pass']} items already meet guidance")
    secs = [x for x in a.section_order if any(r.section == x[0] and r.status == "Pass" for r in a.rows.values())] or a.section_order
    n = max(len(secs), 1)
    w = 9.22 / n
    limit = v2.get("passed_title_chars", 45)
    for k, (disp, pre) in enumerate(secs):
        passed = [r for r in a.rows.values() if r.section == disp and r.status == "Pass"]
        cx = 0.39 + k * w
        tb(s, cx, 1.2, w - 0.1, 0.4, [P(disp, 8.5, True, DK)], name=f"Passed heading {k + 1}")
        tb(s, cx, 1.62, w - 0.1, 0.2, [P(f"{len(passed)} pass", 8, False, GREY)], name=f"Passed count {k + 1}")
        paras = [P([(r.id + "  ", {"bold": True, "color": "0B6B55", "size": 7.5}), (clip(r.title, limit), {})], 8, after=4)
                 for r in sorted(passed, key=lambda r: r.id)]
        tb(s, cx, 1.92, w - 0.1, 3.2, paras, name=f"Passed list {k + 1}")


def detail_slides(d, a, v2):
    order = {"Fail": 0, "Partial": 1, "Needs Review": 2}
    widths = [0.85, 1.75, 0.5, 0.8, 3.3, 2.0]
    top, bottom, hdr = 1.1, 5.0, 0.28
    cap = v2.get("detail_max_rows", 7)
    for disp, _ in a.section_order:
        items = [r for r in a.rows.values() if r.section == disp and r.status in order]
        items.sort(key=lambda r: (order[r.status], r.priority != "Must", r.id))
        if not items:
            continue
        # pack rows by estimated height so long findings get fewer rows per slide
        chunks, cur, used = [], [], hdr
        for r in items:
            lines = max(est_lines(r.title, widths[1] - 0.11, 8.5),
                        1 + est_lines("\n".join(r.evidence), widths[4] - 0.11, 8) + est_lines(r.headline, widths[4] - 0.11, 8.5) - 1,
                        est_lines(r.action or "-", widths[5] - 0.11, 8))
            h = max(0.42, 0.10 + 0.13 * lines)
            if cur and (used + h > bottom - top or len(cur) >= cap):
                chunks.append(cur); cur, used = [], hdr
            cur.append((r, h)); used += h
        if cur:
            chunks.append(cur)
        if len(chunks) > 1:  # balance rows across the slides we need, so no slide is left with a single row
            flat = [x for ch in chunks for x in ch]
            target, chunks, cur, acc = sum(h for _, h in flat) / len(chunks), [], [], 0.0
            for k, (r, h) in enumerate(flat):
                cur.append((r, h)); acc += h
                if acc >= target and len(chunks) < math.ceil(sum(h for _, h in flat) / target) - 1 and k < len(flat) - 1:
                    chunks.append(cur); cur, acc = [], 0.0
            chunks.append(cur)
        c = a.section_counts(disp)
        bits = [f"{c[k]} {k.lower()}" for k in ("Fail", "Partial", "Needs Review") if c[k]]
        for ci, ch in enumerate(chunks):
            suffix = f" ({ci + 1}/{len(chunks)})" if len(chunks) > 1 else ""
            s = d.cream(f"{clip(disp, 40)}: {', '.join(bits)}{suffix}", 18)
            body = []
            for r, _ in ch:
                find = [P(r.headline, 8.5, True, INK, after=1)] + [P(e, 8, False, GREY) for e in r.evidence]
                body.append([[P(r.id, 8, True, DK)], [P(r.title, 8.5)], [P(r.priority, 8, r.priority == "Must")],
                             [P(r.status, 8, True, STATUS[r.status][1])], find, [P(r.action or "-", 8)]])
            tbl = table(s, 0.39, top, widths, ["ID", "Item", "Pri.", "Status", "Finding", "Action"], body,
                        [h for _, h in ch], hdr, name="Detail table")
            for i, (r, _) in enumerate(ch, start=1):
                tbl.cell(i, 3).fill.solid(); tbl.cell(i, 3).fill.fore_color.rgb = RGB(STATUS[r.status][0])
                # re-order fill after borders (schema: borders before fill)
                tcPr = tbl.cell(i, 3)._tc.get_or_add_tcPr()
                for e in [e for e in tcPr if e.tag == qn("a:solidFill")]:
                    tcPr.remove(e); tcPr.append(e)


def na_slide(d, a, narr):
    groups = narr.get("na_groups")
    if groups is None:
        groups = [{"heading": h, "note": n, "ids": ids} for h, n, ids in a.na_groups]
    if not groups:
        return
    s = d.cream(f"Not applicable: {a.na} items that do not apply to this deployment")
    if len(groups) > 6:
        groups = groups[:5] + [{"heading": "Other", "note": "", "ids": [i for g in groups[5:] for i in g["ids"]]}]
    rows_n = math.ceil(len(groups) / 2)
    ch = 1.7 if rows_n <= 2 else 1.1
    for i, g in enumerate(groups):
        cx = 0.39 + (i % 2) * 4.65; cy = 1.25 + (i // 2) * (ch + 0.15)
        card(s, cx, cy, 4.55, ch, name=f"N/A card {i + 1}")
        tb(s, cx + 0.18, cy + 0.12, 3.6, 0.3, [P(g["heading"], 11.5, True, DK)], "m", name=f"N/A heading {i + 1}")
        chip(s, cx + 3.95, cy + 0.15, 0.4, 0.24, str(len(g["ids"])), "Not Applicable", size=9, name=f"N/A count {i + 1}")
        tb(s, cx + 0.18, cy + 0.5, 4.2, 0.3, [P(g.get("note", ""), 9.5, False, GREY)], name=f"N/A note {i + 1}")
        tb(s, cx + 0.18, cy + (0.95 if ch > 1.5 else 0.78), 4.2, 0.6, [P(", ".join(g["ids"]), 8.5)], name=f"N/A items {i + 1}")


# ------------------------------------------------------------------ cover
def edit_cover(slide, data):
    md = data.bundle_metadata
    vals = []
    if md.get("Node / Version"):
        vals.append(f"Node / version: {md['Node / Version']}")
    if md.get("Bundle"):
        vals.append(f"Bundle: {md['Bundle']}")
    if md.get("Report generated"):
        vals.append(f"Report date: {md['Report generated']}")
    for sh in slide.shapes:
        if sh.name == "Bundle metadata":
            paras = list(sh.text_frame.paragraphs)
            for i, para in enumerate(paras):
                if i < len(vals) and para.runs:
                    para.runs[0].text = vals[i]
                    for r in para.runs[1:]:
                        r._r.getparent().remove(r._r)
                else:
                    para._p.getparent().remove(para._p)  # includes the generator-version line


# ------------------------------------------------------------ orchestration
def build_v2(checklist_path: Path, customer: str, output_path: Path, base_deck: Path,
             logo_path: Path | None = None, narrative_path: Path | None = None) -> dict:
    section_config = common.load_config("section_names.yaml")
    layout_config = common.load_config("deck_layout.yaml")
    v2 = layout_config.get("v2", {})
    data = read_checklist.read_checklist(checklist_path, section_config)
    ordered = section_names.order_sections(data.sheet_tab_names, section_config)
    a = da.analyze(data, ordered, v2)
    narr = {}
    warnings = build_deck.check_data_consistency(data, ordered, section_config) + list(a.warnings)
    if not narrative_path:
        found = narr_mod.default_path(checklist_path)
        narrative_path = found if found.exists() else None
        if narrative_path:
            logger.info("Using narrative next to the checklist: %s", narrative_path)
    if narrative_path:
        narr = narr_mod.load(narrative_path)
        warnings += narr_mod.validate(narr, a)
        warnings += narr_mod.staleness_warnings(narr, hashlib.sha256(Path(checklist_path).read_bytes()).hexdigest())

    with tempfile.TemporaryDirectory(prefix="deckv2_") as tmp:
        v1_path = Path(tmp) / "v1.pptx"
        build_deck.build_deck(
            checklist_path=checklist_path, customer=customer, output_path=v1_path, base_deck=base_deck,
            logo_path=logo_path, rows_per_slide=layout_config.get("table_rows_per_slide", 4),
            include_pass_items=False,
        )
        prs = Presentation(str(v1_path))
        old_ids = list(prs.slides._sldIdLst)
        keep_cover, keep_end = old_ids[0], old_ids[-1]
        d = Deck(prs)
        edit_cover(prs.slides[0], data)

        narrative_slides = [(verdict_slide, ()), (snapshot_slide, (data,)), (risk1_slide, ()), (risk2_slide, ()),
                            (risk3_slide, ()), (quickwins_slide, ()), (owners_slide, ()), (roadmap_slide, (data,))]
        for build, extra in narrative_slides:
            try:
                build(d, a, narr, *extra)
            except (TypeError, KeyError, AttributeError) as exc:
                if not narrative_path:
                    raise
                raise narr_mod.NarrativeError(
                    f"{build.__name__.removesuffix('_slide')} slide could not be built from the narrative "
                    f"({type(exc).__name__}: {exc}); check that key's fields against analyze_checklist's `shape`."
                ) from exc
        appendix_divider(d)
        method_slide(d, a, data)
        passed_slide(d, a, v2)
        detail_slides(d, a, v2)
        na_slide(d, a, narr)

        lst = prs.slides._sldIdLst
        new_ids = [e for e in lst if e not in old_ids]
        for e in old_ids:
            if e is not keep_cover and e is not keep_end:
                prs.part.drop_rel(e.rId)
        for e in list(lst):
            lst.remove(e)
        for e in [keep_cover] + new_ids + [keep_end]:
            lst.append(e)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))
        slide_count = len(lst)

    if not narrative_path:
        logger.warning(narr_mod.MISSING_WARNING)
    logger.info("Wrote v2 deck (%d slides): %s", slide_count, output_path)
    return {
        "output_path": output_path, "slide_count": slide_count, "warnings": warnings,
        "narrative_used": str(narrative_path) if narrative_path else None,
        "narrative_missing": not narrative_path,
        "applicable_count": a.applicable, "pass_count": a.counts["Pass"],
        "quick_wins": [q.id for q in a.quick_wins],
        "owner_groups": {o: ids for o, ids in a.owners},
    }
